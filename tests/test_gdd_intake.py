import base64
import json
import re
import shutil
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from helpers import load_script, run_main

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "plugins" / "game-server" / "skills" / "gdd-intake"


convert = load_script(SKILL / "scripts" / "convert_office.py")

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
XML = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def write_zip(path, members):
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def rels(*items):
    body = "".join(
        f'<Relationship Id="{rid}" Type="{kind}" Target="{target}"/>' for rid, kind, target in items
    )
    return f'{XML}<Relationships xmlns="{PKG_REL}">{body}</Relationships>'


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.inputs = self.tmp / "in"
        self.work = self.tmp / "work"
        self.inputs.mkdir()

    def convert(self):
        code, out, err = run_main(convert, [str(self.inputs), str(self.work)])
        self.assertEqual(code, 0, err)
        return json.loads((self.work / "conversion-manifest.json").read_text(encoding="utf-8"))


def make_pptx(path):
    def shape(text, placeholder=None):
        ph = f'<p:nvPr><p:ph type="{placeholder}"/></p:nvPr>' if placeholder else "<p:nvPr/>"
        return (
            f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="s"/><p:cNvSpPr/>{ph}</p:nvSpPr>'
            f"<p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>"
        )

    def slide(*shapes):
        return (
            f'{XML}<p:sld xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:cSld><p:spTree>'
            + "".join(shapes)
            + "</p:spTree></p:cSld></p:sld>"
        )

    table = (
        "<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
        "<a:tr><a:tc><a:txBody><a:p><a:r><a:t>Tier</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>Points</a:t></a:r></a:p></a:txBody></a:tc></a:tr>"
        "<a:tr><a:tc><a:txBody><a:p><a:r><a:t>1</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>100</a:t></a:r></a:p></a:txBody></a:tc></a:tr>"
        "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    )
    picture = '<p:pic><p:blipFill><a:blip r:embed="rId7"/></p:blipFill></p:pic>'
    notes = (
        f'{XML}<p:notes xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree>'
        + shape("1", "sldNum")
        + shape("Say the points slowly.", "body")
        + "</p:spTree></p:cSld></p:notes>"
    )
    presentation = (
        f'{XML}<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>'
        '<p:sldId id="256" r:id="rId1"/><p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>'
    )
    write_zip(
        path,
        {
            "[Content_Types].xml": f"{XML}<Types/>",
            "ppt/presentation.xml": presentation,
            "ppt/_rels/presentation.xml.rels": rels(
                ("rId1", f"{R}/slide", "slides/slide2.xml"),
                ("rId2", f"{R}/slide", "slides/slide1.xml"),
            ),
            "ppt/slides/slide2.xml": slide(shape("Introduction", "title"), shape("Earn points daily.")),
            "ppt/slides/slide1.xml": slide(shape("Rewards", "ctrTitle"), table, picture),
            "ppt/slides/_rels/slide1.xml.rels": rels(
                ("rId7", f"{R}/image", "../media/image1.png"),
                ("rId8", f"{R}/notesSlide", "../notesSlides/notesSlide1.xml"),
            ),
            "ppt/notesSlides/notesSlide1.xml": notes,
            "ppt/media/image1.png": PNG,
        },
    )


C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def make_pptx_with_chart(path):
    def slide(title, extra=""):
        return (
            f'{XML}<p:sld xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:cSld><p:spTree>'
            f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="t"/><p:cNvSpPr/><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>'
            f"<p:txBody><a:bodyPr/><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp>{extra}"
            "</p:spTree></p:cSld></p:sld>"
        )

    frame = (
        f'<p:graphicFrame><a:graphic><a:graphicData uri="{C}">'
        f'<c:chart xmlns:c="{C}" r:id="rId3"/></a:graphicData></a:graphic></p:graphicFrame>'
    )

    def series(name, values):
        points = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values))
        categories = "".join(f'<c:pt idx="{i}"><c:v>Week {i + 1}</c:v></c:pt>' for i in range(len(values)))
        return (
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt idx=\"0\"><c:v>{name}</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache>{categories}</c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache>{points}</c:numCache></c:numRef></c:val></c:ser>"
        )

    chart = (
        f'{XML}<c:chartSpace xmlns:c="{C}" xmlns:a="{A}"><c:chart>'
        "<c:title><c:tx><c:rich><a:p><a:r><a:t>Weekly players</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        "<c:plotArea><c:barChart>"
        + series("Free", ["120", "0.28999999999999998"])
        + series("Paid", ["30", "45"])
        + "</c:barChart></c:plotArea></c:chart></c:chartSpace>"
    )
    presentation = (
        f'{XML}<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>'
        '<p:sldId id="256" r:id="rId1"/><p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>'
    )
    write_zip(
        path,
        {
            "[Content_Types].xml": f"{XML}<Types/>",
            "ppt/presentation.xml": presentation,
            "ppt/_rels/presentation.xml.rels": rels(
                ("rId1", f"{R}/slide", "slides/slide1.xml"),
                ("rId2", f"{R}/slide", "slides/slide2.xml"),
            ),
            "ppt/slides/slide1.xml": slide("Players", frame),
            "ppt/slides/_rels/slide1.xml.rels": rels(("rId3", f"{R}/chart", "../charts/chart1.xml")),
            "ppt/charts/chart1.xml": chart,
            "ppt/slides/slide2.xml": slide("Next"),
        },
    )


class ConvertTest(TempDirCase):
    def test_pptx_follows_deck_order_and_keeps_tables_notes_and_pictures(self):
        make_pptx(self.inputs / "deck.pptx")
        self.convert()
        text = (self.work / "raw" / "slides_deck.md").read_text(encoding="utf-8")
        first = text.index("## Slide 1: Introduction")
        second = text.index("## Slide 2: Rewards")
        self.assertLess(first, second)
        self.assertIn("Earn points daily.", text)
        self.assertIn("| Tier | Points |", text)
        self.assertIn("| 1 | 100 |", text)
        self.assertIn("Say the points slowly.", text[second:])
        self.assertNotIn("Say the points slowly.", text[:second])
        self.assertEqual([p.read_bytes() for p in (self.work / "images").iterdir()], [PNG])

    def test_loose_images_that_share_a_stem_get_their_own_names(self):
        (self.inputs / "a.jpeg").write_bytes(b"jpeg bytes")
        (self.inputs / "a.jpg").write_bytes(b"jpg bytes")
        manifest = self.convert()
        outputs = {f["source"]: f["output"] for f in manifest["files"]}
        self.assertEqual(outputs, {"a.jpeg": "images/standalone-a.jpg", "a.jpg": "images/standalone-a_.jpg"})
        for source, output in outputs.items():
            self.assertEqual((self.work / output).read_bytes(), (self.inputs / source).read_bytes())

    def test_text_files_pass_through_and_legacy_formats_are_skipped(self):
        (self.inputs / "notes.md").write_text("# Notes\n", encoding="utf-8")
        (self.inputs / "old.ppt").write_bytes(b"x")
        manifest = self.convert()
        self.assertEqual((self.work / "raw" / "notes.md").read_text(encoding="utf-8"), "# Notes\n")
        skipped = {item["file"]: item["reason"] for item in manifest["skipped"]}
        self.assertIn("pptx", skipped["old.ppt"])

    def test_pdf_without_pymupdf_asks_for_an_export(self):
        if convert.HAS_PYMUPDF:
            self.skipTest("PyMuPDF is installed")
        (self.inputs / "spec.pdf").write_bytes(b"%PDF-1.4\n")
        code, out, err = run_main(convert, [str(self.inputs), str(self.work)])
        self.assertEqual(code, 0)
        manifest = json.loads((self.work / "conversion-manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([r["file"] for r in manifest["export_requests"]], ["spec.pdf"])
        self.assertIn("spec.pdf", out)
        self.assertIn("docx or Markdown", out)
        self.assertEqual(manifest["files"], [])

    def test_a_broken_file_is_reported_without_stopping_the_rest(self):
        (self.inputs / "broken.pptx").write_bytes(b"not a zip")
        (self.inputs / "notes.txt").write_text("hello", encoding="utf-8")
        manifest = self.convert()
        self.assertEqual([w["file"] for w in manifest["warnings"]], ["broken.pptx"])
        self.assertTrue((self.work / "raw" / "notes.txt").exists())

    def test_missing_input_folder_fails(self):
        code, _, _ = run_main(convert, [str(self.tmp / "absent"), str(self.work)])
        self.assertEqual(code, 1)

    def test_pptx_chart_title_categories_and_series_land_under_their_slide(self):
        make_pptx_with_chart(self.inputs / "deck.pptx")
        self.convert()
        text = (self.work / "raw" / "slides_deck.md").read_text(encoding="utf-8")
        slide = text.split("## Slide 1: Players", 1)[1].split("## Slide 2", 1)[0]
        self.assertIn("Weekly players", slide)
        self.assertIn("| Category | Free | Paid |", slide)
        self.assertIn("| Week 1 | 120 | 30 |", slide)
        self.assertIn("| Week 2 | 0.29 | 45 |", slide)

    def test_pdf_text_is_nfkc_normalized(self):
        class Page:
            def get_text(self, kind):
                return "Proﬁt per ﬁsh"

            def get_images(self, full):
                return []

        class Document:
            page_count = 1

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def __iter__(self):
                return iter([Page()])

        fake = types.SimpleNamespace(open=lambda path: Document())
        (self.inputs / "spec.pdf").write_bytes(b"%PDF-1.4\n")
        with mock.patch.multiple(convert, fitz=fake, HAS_PYMUPDF=True):
            self.convert()
        text = (self.work / "raw" / "pdf_spec.md").read_text(encoding="utf-8")
        self.assertIn("Profit per fish", text)

    def test_docx_and_xlsx_are_skipped_with_the_install_command_without_the_packages(self):
        (self.inputs / "challenge.docx").write_bytes(b"x")
        (self.inputs / "balance.xlsx").write_bytes(b"x")
        make_pptx(self.inputs / "deck.pptx")
        with mock.patch.object(convert, "HAS_GDDLIB", False):
            manifest = self.convert()
        skipped = {item["file"]: item["reason"] for item in manifest["skipped"]}
        self.assertEqual(sorted(skipped), ["balance.xlsx", "challenge.docx"])
        self.assertIn("pip install python-docx openpyxl", skipped["challenge.docx"])
        self.assertEqual([f["source"] for f in manifest["files"]], ["deck.pptx"])


@unittest.skipUnless(convert.HAS_GDDLIB, "python-docx and openpyxl are not installed")
class GddlibConvertTest(TempDirCase):
    def make_docx(self, name):
        from docx import Document

        document = Document()
        document.add_heading("Weekly challenge", level=1)
        paragraph = document.add_paragraph("Day one gives ")
        paragraph.add_run("500 gold").bold = True
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip(table._cells, ["Day", "Gold", "1", "500"]):
            cell.text = value
        image = self.tmp / "pixel.png"
        image.write_bytes(PNG)
        document.add_picture(str(image))
        document.save(self.inputs / name)

    def make_xlsx(self, name):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rewards"
        sheet.append(["Day", "Gold", "Double"])
        sheet.append([1, 500, "=B2*2"])
        sheet.merge_cells("A4:B4")
        sheet["A4"] = "Week one"
        workbook.save(self.inputs / name)

    def test_docx_becomes_markdown_with_its_images_beside_it(self):
        self.make_docx("challenge.docx")
        manifest = self.convert()
        entry = manifest["files"][0]
        self.assertEqual((entry["type"], entry["output"]), ("docx", "raw/challenge.md"))
        text = (self.work / "raw" / "challenge.md").read_text(encoding="utf-8")
        self.assertIn("# Weekly challenge", text)
        self.assertIn("**500 gold**", text)
        self.assertIn("| Day | Gold |", text)
        self.assertEqual(entry["images_extracted"], 1)
        linked = re.findall(r"\]\((assets/[^)]+)\)", text)
        self.assertEqual(len(linked), 1, text)
        self.assertEqual((self.work / "raw" / linked[0]).read_bytes(), PNG)

    def test_xlsx_becomes_one_markdown_file_with_a_section_per_sheet(self):
        self.make_xlsx("balance.xlsx")
        manifest = self.convert()
        entry = manifest["files"][0]
        self.assertEqual((entry["type"], entry["output"]), ("xlsx", "raw/balance.md"))
        self.assertEqual(entry["sheets"][0]["sheet"], "Rewards")
        text = (self.work / "raw" / "balance.md").read_text(encoding="utf-8")
        self.assertIn("## Sheet: Rewards", text)
        self.assertIn("| Day | Gold | Double |", text)
        self.assertIn("| 1 | 500 | =B2*2 |", text)
        self.assertIn("| Week one | Week one |  |", text)
        self.assertFalse(list((self.work / "raw").glob("*.csv")))

    def test_a_docx_and_an_xlsx_with_one_name_report_the_second(self):
        self.make_docx("challenge.docx")
        self.make_xlsx("challenge.xlsx")
        manifest = self.convert()
        self.assertEqual([f["source"] for f in manifest["files"]], ["challenge.docx"])
        self.assertEqual([w["file"] for w in manifest["warnings"]], ["challenge.xlsx"])
        self.assertIn("rename", manifest["warnings"][0]["error"])

    def test_a_docx_that_fails_leaves_its_name_to_an_xlsx(self):
        (self.inputs / "balance.docx").write_bytes(b"not a zip")
        self.make_xlsx("balance.xlsx")
        manifest = self.convert()
        self.assertEqual([f["output"] for f in manifest["files"]], ["raw/balance.md"])
        self.assertEqual([w["file"] for w in manifest["warnings"]], ["balance.docx"])

    def test_a_symlinked_workbook_converts(self):
        shared = self.tmp / "shared"
        shared.mkdir()
        self.make_xlsx("balance.xlsx")
        (self.inputs / "balance.xlsx").rename(shared / "balance.xlsx")
        (self.inputs / "balance.xlsx").symlink_to(shared / "balance.xlsx")
        manifest = self.convert()
        self.assertEqual(manifest["warnings"], [])
        self.assertIn("## Sheet: Rewards", (self.work / "raw" / "balance.md").read_text(encoding="utf-8"))

    def test_a_text_file_named_like_a_converted_file_gets_its_own_name(self):
        self.make_docx("a.docx")
        (self.inputs / "A.md").write_text("# Notes\n", encoding="utf-8")
        manifest = self.convert()
        outputs = {f["source"]: f["output"] for f in manifest["files"]}
        self.assertEqual(outputs, {"a.docx": "raw/a.md", "A.md": "raw/A_.md"})
        self.assertEqual((self.work / "raw" / "A_.md").read_text(encoding="utf-8"), "# Notes\n")
        self.assertIn("Weekly challenge", (self.work / "raw" / "a.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
