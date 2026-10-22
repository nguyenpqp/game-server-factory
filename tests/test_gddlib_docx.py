from __future__ import annotations

import importlib.util
import sys
import unittest
from pathlib import Path

if importlib.util.find_spec("docx") is None:
    raise unittest.SkipTest("python-docx is not installed")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "plugins" / "game-server" / "skills" / "gdd-intake" / "scripts"))

import base64
from collections.abc import Callable
from copy import deepcopy
import hashlib
from io import BytesIO
from pathlib import Path
import shutil
import unittest
from uuid import uuid4
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from gddlib.docx import convert_docx_file
from gddlib.markdown import verify_docx_markdown


class DocxMarkdownConversionTest(unittest.TestCase):
    TINY_PNG = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
        "AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )

    def test_adjacent_runs_with_the_same_style_share_one_markdown_span(
        self,
    ) -> None:
        cases = (
            ((True, False, False), "**Hộp**"),
            ((False, True, False), "*Hộp*"),
            ((False, False, True), "~~Hộp~~"),
            ((True, True, False), "***Hộp***"),
            ((True, False, True), "**~~Hộp~~**"),
            ((False, True, True), "*~~Hộp~~*"),
            ((True, True, True), "***~~Hộp~~***"),
        )

        def build(document: Document) -> None:
            for (bold, italic, strike), _ in cases:
                paragraph = document.add_paragraph()
                for text in ("H", "ộ", "p"):
                    run = paragraph.add_run(text)
                    run.bold = bold
                    run.italic = italic
                    run.font.strike = strike

        markdown = self._convert(build)

        for _, expected in cases:
            with self.subTest(expected=expected):
                self.assertIn(expected, markdown)
        self.assertNotIn("****", markdown)
        self.assertNotIn("~~~~", markdown)

    def test_word_toc_links_target_markdown_heading_fragments(self) -> None:
        def build(document: Document) -> None:
            heading = document.add_paragraph(style="Heading 1")
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "1")
            bookmark_start.set(qn("w:name"), "_Toc123")
            heading._p.append(bookmark_start)
            heading.add_run("1. Hộp Thư")
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), "1")
            heading._p.append(bookmark_end)

            toc_entry = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc123")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1. Hộp Thư"
            run.append(text)
            hyperlink.append(run)
            toc_entry._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn("# 1\\. Hộp Thư", markdown)
        self.assertIn(
            "[1\\. Hộp Thư](<#1-hộp-thư>)",
            markdown,
        )
        self.assertNotIn('<a id="_Toc123"></a>', markdown)
        self.assertNotIn("(<#_Toc123>)", markdown)

    def test_toc_style_entries_are_rebuilt_from_actual_headings(self) -> None:
        def build(document: Document) -> None:
            toc_style = document.styles.add_style(
                "toc 1",
                WD_STYLE_TYPE.PARAGRAPH,
            )
            document.add_paragraph("Mục Lục")
            document.add_paragraph("1. Tổng Quan\t2", style=toc_style)
            document.add_paragraph(
                "3. Flow Chính — Kết Bạn\t2",
                style=toc_style,
            )
            document.add_paragraph("6. Phần Thưởng\t8", style=toc_style)
            document.add_paragraph(
                "7. Cửa Hàng Ven Đường & Mở Ô Nhờ Bạn Bè\t9",
                style=toc_style,
            )
            document.add_heading("1. Tổng Quan", level=1)
            document.add_heading("3. Flow Chính", level=1)
            document.add_heading(
                "6. Cửa Hàng Ven Đường & Mở Ô Nhờ Bạn Bè",
                level=1,
            )

        markdown = self._convert(build)

        self.assertIn("[1\\. Tổng Quan](<#1-tổng-quan>)", markdown)
        self.assertIn("[3\\. Flow Chính](<#3-flow-chính>)", markdown)
        self.assertIn(
            "[6\\. Cửa Hàng Ven Đường \\& Mở Ô Nhờ Bạn Bè]"
            "(<#6-cửa-hàng-ven-đường--mở-ô-nhờ-bạn-bè>)",
            markdown,
        )
        self.assertNotIn("Phần Thưởng", markdown)
        self.assertNotIn("7\\. Cửa Hàng", markdown)

    def test_verifier_checks_fragment_only_links(self) -> None:
        temporary_root = (
            Path(__file__).resolve().parents[5]
            / "build"
            / "gdd-sync"
            / "test-tmp"
        )
        temporary_root.mkdir(parents=True, exist_ok=True)
        root = temporary_root / uuid4().hex
        root.mkdir()
        try:
            markdown = root / "doc.md"
            markdown.write_text(
                "# Heading thật\n\n"
                "[Đúng](#heading-thật)\n\n"
                "[Sai](#heading-không-tồn-tại)\n",
                encoding="utf-8",
                newline="\n",
            )

            summary, errors = verify_docx_markdown(markdown, root)

            self.assertEqual(1, summary["broken_links"])
            self.assertTrue(
                any("missing local link fragment" in error for error in errors)
            )
        finally:
            shutil.rmtree(root)

    def test_verifier_rejects_docx_markdown_artifacts(self) -> None:
        temporary_root = (
            Path(__file__).resolve().parents[5]
            / "build"
            / "gdd-sync"
            / "test-tmp"
        )
        temporary_root.mkdir(parents=True, exist_ok=True)
        root = temporary_root / uuid4().hex
        root.mkdir()
        try:
            markdown = root / "doc.md"
            markdown.write_text(
                "# Heading\n\n"
                '<a id="_Toc123"></a>\n\n'
                "Nội dung **** lỗi.\n",
                encoding="utf-8",
                newline="\n",
            )

            summary, errors = verify_docx_markdown(markdown, root)

            self.assertEqual(2, summary["errors"])
            self.assertTrue(
                any("raw Word TOC bookmark" in error for error in errors)
            )
            self.assertTrue(
                any("degenerate style delimiter" in error for error in errors)
            )
        finally:
            shutil.rmtree(root)

    def test_word_toc_slug_preserves_spaces_around_punctuation(
        self,
    ) -> None:
        def build(document: Document) -> None:
            heading = document.add_paragraph(style="Heading 1")
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "2")
            bookmark_start.set(qn("w:name"), "_Toc124")
            heading._p.append(bookmark_start)
            heading.add_run(
                "2. Cơ Chế Hiển Thị & Điểm Truy Cập"
            )
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), "2")
            heading._p.append(bookmark_end)

            toc_entry = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc124")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Section"
            run.append(text)
            hyperlink.append(run)
            toc_entry._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn(
            "[Section](<#2-cơ-chế-hiển-thị--điểm-truy-cập>)",
            markdown,
        )

    def test_word_toc_slug_uses_only_visible_heading_text(self) -> None:
        def build(document: Document) -> None:
            heading = document.add_paragraph(style="Heading 1")
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "3")
            bookmark_start.set(qn("w:name"), "_Toc125")
            heading._p.append(bookmark_start)
            hidden = heading.add_run("Secret ")
            hidden.font.hidden = True
            heading.add_run("Title")

            source = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc125")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Visible"
            run.append(text)
            hyperlink.append(run)
            source._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn("# Title", markdown)
        self.assertIn("[Visible](<#title>)", markdown)
        self.assertNotIn("secret-title", markdown)

    def test_word_toc_slug_matches_github_unicode_lowercasing(
        self,
    ) -> None:
        def build(document: Document) -> None:
            heading = document.add_paragraph(style="Heading 1")
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "4")
            bookmark_start.set(qn("w:name"), "_Toc126")
            heading._p.append(bookmark_start)
            heading.add_run("Straße")

            source = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc126")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Unicode"
            run.append(text)
            hyperlink.append(run)
            source._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn("[Unicode](<#straße>)", markdown)
        self.assertNotIn("<#strasse>", markdown)

    def test_word_toc_slug_avoids_existing_suffixed_collision(
        self,
    ) -> None:
        def add_heading(
            document: Document,
            text: str,
            bookmark_id: int,
        ) -> None:
            heading = document.add_paragraph(style="Heading 1")
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), str(bookmark_id))
            bookmark_start.set(
                qn("w:name"),
                f"_Toc{200 + bookmark_id}",
            )
            heading._p.append(bookmark_start)
            heading.add_run(text)

        def build(document: Document) -> None:
            add_heading(document, "foo", 1)
            add_heading(document, "foo-1", 2)
            add_heading(document, "foo", 3)

            source = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc203")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Third"
            run.append(text)
            hyperlink.append(run)
            source._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn("[Third](<#foo-2>)", markdown)

    def test_style_markers_exclude_trailing_run_whitespace(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            bold = paragraph.add_run("Bold ")
            bold.bold = True
            paragraph.add_run("plain")

        markdown = self._convert(build)

        self.assertIn("**Bold** plain", markdown)
        self.assertNotIn("**Bold **", markdown)

    def test_same_style_runs_merge_across_unstyled_whitespace(
        self,
    ) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            first = paragraph.add_run("Cơ Chế ")
            first.bold = True
            second = paragraph.add_run("Hiển Thị")
            second.bold = True
            paragraph.add_run(" ")
            third = paragraph.add_run("& Điểm Truy Cập")
            third.bold = True

        markdown = self._convert(build)

        self.assertIn(
            "**Cơ Chế Hiển Thị \\& Điểm Truy Cập**",
            markdown,
        )
        self.assertNotIn("** **", markdown)

    def test_shared_style_stays_open_when_the_next_run_adds_style(
        self,
    ) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            first = paragraph.add_run("A")
            first.bold = True
            second = paragraph.add_run("B")
            second.bold = True
            second.font.strike = True

        markdown = self._convert(build)

        self.assertIn("**A~~B~~**", markdown)
        self.assertNotIn("****", markdown)

    def test_transparent_word_container_does_not_split_style(
        self,
    ) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            first = paragraph.add_run("H")
            first.bold = True

            smart_tag = OxmlElement("w:smartTag")
            nested_run = OxmlElement("w:r")
            run_properties = OxmlElement("w:rPr")
            run_properties.append(OxmlElement("w:b"))
            nested_run.append(run_properties)
            nested_text = OxmlElement("w:t")
            nested_text.text = "ộ"
            nested_run.append(nested_text)
            smart_tag.append(nested_run)
            paragraph._p.append(smart_tag)

            last = paragraph.add_run("p")
            last.bold = True

        markdown = self._convert(build)

        self.assertIn("**Hộp**", markdown)
        self.assertNotIn("****", markdown)

    def test_author_bookmarks_and_links_remain_explicit(self) -> None:
        def build(document: Document) -> None:
            target = document.add_paragraph()
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "7")
            bookmark_start.set(qn("w:name"), "details")
            target._p.append(bookmark_start)
            target.add_run("Details")
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), "7")
            target._p.append(bookmark_end)

            source = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "details")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Jump"
            run.append(text)
            hyperlink.append(run)
            source._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn('<a id="details"></a>Details', markdown)
        self.assertIn("[Jump](<#details>)", markdown)

    def test_hyperlink_preserves_surrounding_label_whitespace(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            paragraph.add_run("Before")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "details")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.set(qn("xml:space"), "preserve")
            text.text = " link "
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
            paragraph.add_run("after")

        markdown = self._convert(build)

        self.assertIn("Before [link](<#details>) after", markdown)

    def test_simple_hyperlink_field_is_preserved(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            field = OxmlElement("w:fldSimple")
            field.set(
                qn("w:instr"),
                ' HYPERLINK "https://example.com/docs" ',
            )
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Docs"
            run.append(text)
            field.append(run)
            paragraph._p.append(field)

        markdown = self._convert(build)

        self.assertIn("[Docs](<https://example.com/docs>)", markdown)

    def test_complex_hyperlink_field_is_preserved(self) -> None:
        def add_field_run(
            paragraph: object,
            field_type: str,
        ) -> None:
            run = OxmlElement("w:r")
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run.append(field)
            paragraph._p.append(run)

        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            add_field_run(paragraph, "begin")
            instruction_run = OxmlElement("w:r")
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = (
                ' HYPERLINK "https://example.com/complex" '
            )
            instruction_run.append(instruction)
            paragraph._p.append(instruction_run)
            add_field_run(paragraph, "separate")
            paragraph.add_run("Complex")
            add_field_run(paragraph, "end")

        markdown = self._convert(build)

        self.assertIn(
            "[Complex](<https://example.com/complex>)",
            markdown,
        )

    def test_complex_hyperlink_instruction_can_share_begin_run(
        self,
    ) -> None:
        def add_field_run(
            paragraph: object,
            field_type: str,
        ) -> None:
            run = OxmlElement("w:r")
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run.append(field)
            paragraph._p.append(run)

        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            begin_run = OxmlElement("w:r")
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText")
            instruction.text = (
                ' HYPERLINK "https://example.com/same-run" '
            )
            begin_run.append(begin)
            begin_run.append(instruction)
            paragraph._p.append(begin_run)
            add_field_run(paragraph, "separate")
            paragraph.add_run("Same run")
            add_field_run(paragraph, "end")

        markdown = self._convert(build)

        self.assertIn(
            "[Same run](<https://example.com/same-run>)",
            markdown,
        )

    def test_ordered_list_syntax_split_across_runs_is_escaped(
        self,
    ) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            paragraph.add_run("1")
            paragraph.add_run(". plain paragraph")

        markdown = self._convert(build)

        self.assertIn("1\\. plain paragraph", markdown)
        self.assertNotIn("\n\n1. plain paragraph", markdown)

    def test_block_container_preserves_nested_paragraph_order(
        self,
    ) -> None:
        def build(document: Document) -> None:
            document.add_paragraph("Before")
            nested = document.add_paragraph("Nested")
            document.add_paragraph("After")

            body = document.element.body
            position = body.index(nested._p)
            body.remove(nested._p)
            container = OxmlElement("w:customXml")
            container.append(nested._p)
            body.insert(position, container)

        markdown = self._convert(build)

        self.assertIn("Before\n\nNested\n\nAfter", markdown)

    def test_inline_move_destination_preserves_visible_text(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("Moved")
            paragraph._p.remove(run._r)
            move_to = OxmlElement("w:moveTo")
            move_to.append(run._r)
            paragraph._p.append(move_to)

        markdown = self._convert(build)

        self.assertIn("\n\nMoved\n", markdown)

    def test_no_break_hyphen_is_not_dropped(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = OxmlElement("w:r")
            before = OxmlElement("w:t")
            before.text = "non"
            run.append(before)
            run.append(OxmlElement("w:noBreakHyphen"))
            after = OxmlElement("w:t")
            after.text = "breaking"
            run.append(after)
            paragraph._p.append(run)

        markdown = self._convert(build)

        self.assertIn("non‑breaking", markdown)

    def test_hidden_run_text_is_omitted(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph("Shown")
            hidden = paragraph.add_run("SECRET")
            hidden.font.hidden = True

        markdown = self._convert(build)

        self.assertIn("\n\nShown\n", markdown)
        self.assertNotIn("SECRET", markdown)

    def test_outline_level_style_renders_heading_and_toc_target(
        self,
    ) -> None:
        def build(document: Document) -> None:
            style = document.styles.add_style(
                "Custom section",
                WD_STYLE_TYPE.PARAGRAPH,
            )
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            style.element.get_or_add_pPr().append(outline)

            heading = document.add_paragraph(style=style)
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), "11")
            bookmark_start.set(qn("w:name"), "_Toc456")
            heading._p.append(bookmark_start)
            heading.add_run("Custom heading")
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), "11")
            heading._p.append(bookmark_end)

            source = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_Toc456")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Go"
            run.append(text)
            hyperlink.append(run)
            source._p.append(hyperlink)

        markdown = self._convert(build)

        self.assertIn("# Custom heading", markdown)
        self.assertIn("[Go](<#custom-heading>)", markdown)

    def test_alternate_content_uses_one_image_representation(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(BytesIO(self.TINY_PNG))
            drawing = run._r.find(qn("w:drawing"))
            assert drawing is not None
            run._r.remove(drawing)

            compatibility_namespace = (
                "http://schemas.openxmlformats.org/"
                "markup-compatibility/2006"
            )
            alternate = etree.Element(
                f"{{{compatibility_namespace}}}AlternateContent"
            )
            choice = etree.Element(
                f"{{{compatibility_namespace}}}Choice"
            )
            choice.set("Requires", "wps")
            choice.append(deepcopy(drawing))
            fallback = etree.Element(
                f"{{{compatibility_namespace}}}Fallback"
            )
            fallback.append(deepcopy(drawing))
            alternate.append(choice)
            alternate.append(fallback)
            run._r.append(alternate)

        markdown = self._convert(build)

        self.assertEqual(1, markdown.count("!["))

    def test_alternate_content_uses_fallback_for_unknown_choice(
        self,
    ) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(BytesIO(self.TINY_PNG))
            drawing = run._r.find(qn("w:drawing"))
            assert drawing is not None
            blip = drawing.find(".//" + qn("a:blip"))
            assert blip is not None
            relationship_id = blip.get(qn("r:embed"))
            assert relationship_id is not None
            run._r.remove(drawing)

            compatibility_namespace = (
                "http://schemas.openxmlformats.org/"
                "markup-compatibility/2006"
            )
            alternate = etree.Element(
                f"{{{compatibility_namespace}}}AlternateContent"
            )
            choice = etree.Element(
                f"{{{compatibility_namespace}}}Choice"
            )
            choice.set("Requires", "wps")
            choice.append(OxmlElement("w:drawing"))
            fallback = etree.Element(
                f"{{{compatibility_namespace}}}Fallback"
            )
            pict = OxmlElement("w:pict")
            vml_namespace = "urn:schemas-microsoft-com:vml"
            shape = etree.Element(f"{{{vml_namespace}}}shape")
            image_data = etree.Element(
                f"{{{vml_namespace}}}imagedata"
            )
            image_data.set(qn("r:id"), relationship_id)
            shape.append(image_data)
            pict.append(shape)
            fallback.append(pict)
            alternate.append(choice)
            alternate.append(fallback)
            run._r.append(alternate)

        markdown = self._convert(build)

        self.assertEqual(1, markdown.count("!["))

    def test_text_around_an_image_keeps_run_style(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("Label ")
            run.bold = True
            run.add_picture(BytesIO(self.TINY_PNG))
            run.add_text(" caption")

        markdown = self._convert(build)

        self.assertIn("**Label** ![", markdown)
        self.assertIn(") **caption**", markdown)

    def test_table_image_alt_is_single_line_and_not_double_escaped(
        self,
    ) -> None:
        def build(document: Document) -> None:
            cell = document.add_table(rows=1, cols=1).cell(0, 0)
            run = cell.paragraphs[0].add_run()
            run.add_picture(BytesIO(self.TINY_PNG))
            drawing = run._r.find(qn("w:drawing"))
            assert drawing is not None
            doc_properties = drawing.find(".//" + qn("wp:docPr"))
            assert doc_properties is not None
            doc_properties.set("descr", "A [B]\nC")

        markdown = self._convert(build)

        self.assertIn("![A &#91;B&#93; C](", markdown)
        self.assertNotIn("\\\\[B", markdown)

    def test_character_style_formatting_is_preserved(self) -> None:
        def build(document: Document) -> None:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("Strong")
            run.style = "Strong"

        markdown = self._convert(build)

        self.assertIn("**Strong**", markdown)

    def test_merged_table_cell_content_is_emitted_once(self) -> None:
        def build(document: Document) -> None:
            table = document.add_table(rows=2, cols=2)
            merged = table.cell(0, 0).merge(table.cell(0, 1))
            merged.text = "Merged"
            table.cell(1, 0).text = "Left"
            table.cell(1, 1).text = "Right"

        markdown = self._convert(build)

        self.assertIn("| Merged |  |", markdown)
        self.assertNotIn("| Merged | Merged |", markdown)

    def test_distinct_table_cells_are_never_deduplicated(self) -> None:
        def build(document: Document) -> None:
            table = document.add_table(rows=8, cols=3)
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    cell.text = f"R{row_index}C{column_index}"

        markdown = self._convert(build)

        for row_index in range(8):
            for column_index in range(3):
                with self.subTest(
                    row=row_index,
                    column=column_index,
                ):
                    self.assertIn(
                        f"R{row_index}C{column_index}",
                        markdown,
                    )

    def test_nested_merged_table_cell_content_is_emitted_once(self) -> None:
        def build(document: Document) -> None:
            outer = document.add_table(rows=1, cols=1)
            nested = outer.cell(0, 0).add_table(rows=1, cols=2)
            merged = nested.cell(0, 0).merge(nested.cell(0, 1))
            merged.text = "Nested merged"

        markdown = self._convert(build)

        self.assertIn("Nested merged", markdown)
        self.assertNotIn("Nested merged / Nested merged", markdown)

    def test_same_version_legacy_assets_are_migrated_safely(self) -> None:
        root, source_root, output_root, asset_dir = (
            self._legacy_asset_fixture(shared=False)
        )
        try:
            result = convert_docx_file(
                source_root,
                "Mailbox.docx",
                output_root,
            )

            expected_name = (
                "image_1_"
                f"{hashlib.sha256(self.TINY_PNG).hexdigest()[:12]}.png"
            )
            self.assertEqual("converted", result.status)
            self.assertEqual([expected_name], sorted(
                path.name for path in asset_dir.iterdir()
            ))
            markdown = (output_root / "Mailbox.md").read_text(
                encoding="utf-8"
            )
            self.assertIn(f"/{expected_name})", markdown)
        finally:
            shutil.rmtree(root)

    def test_root_level_media_relationship_is_converted(self) -> None:
        temporary_root = (
            Path(__file__).resolve().parents[5]
            / "build"
            / "gdd-sync"
            / "test-tmp"
        )
        temporary_root.mkdir(parents=True, exist_ok=True)
        root = temporary_root / uuid4().hex
        source_root = root / "source"
        output_root = root / "output"
        source_root.mkdir(parents=True)
        output_root.mkdir(parents=True)
        source = source_root / "Mailbox.docx"
        image = root / "image.png"
        image.write_bytes(self.TINY_PNG)
        document = Document()
        document.add_picture(str(image))
        document.save(source)
        patched = root / "patched.docx"
        try:
            with zipfile.ZipFile(source) as old_archive:
                with zipfile.ZipFile(patched, "w") as new_archive:
                    for info in old_archive.infolist():
                        payload = old_archive.read(info.filename)
                        if info.filename == "word/_rels/document.xml.rels":
                            updated = payload.replace(
                                b'Target="media/image1.png"',
                                b'Target="../media/image1.png"',
                            )
                            self.assertNotEqual(payload, updated)
                            payload = updated
                        new_archive.writestr(info, payload)
                    new_archive.writestr(
                        "media/image1.png",
                        self.TINY_PNG,
                    )
            patched.replace(source)

            result = convert_docx_file(
                source_root,
                "Mailbox.docx",
                output_root,
            )

            expected_name = (
                "image_1_"
                f"{hashlib.sha256(self.TINY_PNG).hexdigest()[:12]}.png"
            )
            markdown = (output_root / "Mailbox.md").read_text(
                encoding="utf-8"
            )
            self.assertEqual("converted", result.status)
            self.assertIn(f"/{expected_name})", markdown)
        finally:
            shutil.rmtree(root)

    def test_same_version_assets_shared_by_another_markdown_are_rejected(
        self,
    ) -> None:
        root, source_root, output_root, asset_dir = (
            self._legacy_asset_fixture(shared=True)
        )
        try:
            with self.assertRaisesRegex(
                RuntimeError,
                "old asset directory is still referenced",
            ):
                convert_docx_file(
                    source_root,
                    "Mailbox.docx",
                    output_root,
                )

            self.assertTrue((asset_dir / "image1.png").is_file())
            self.assertEqual(
                ["image1.png"],
                sorted(path.name for path in asset_dir.iterdir()),
            )
        finally:
            shutil.rmtree(root)

    def _legacy_asset_fixture(
        self,
        *,
        shared: bool,
    ) -> tuple[Path, Path, Path, Path]:
        temporary_root = (
            Path(__file__).resolve().parents[5]
            / "build"
            / "gdd-sync"
            / "test-tmp"
        )
        temporary_root.mkdir(parents=True, exist_ok=True)
        root = temporary_root / uuid4().hex
        source_root = root / "source"
        output_root = root / "output"
        source_root.mkdir(parents=True)
        output_root.mkdir(parents=True)
        source = source_root / "Mailbox.docx"
        image = root / "image.png"
        image.write_bytes(self.TINY_PNG)
        document = Document()
        document.add_picture(str(image))
        document.save(source)

        source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
        asset_dir = output_root / "assets" / f"mailbox-{source_sha[:12]}"
        asset_dir.mkdir(parents=True)
        (asset_dir / "image1.png").write_bytes(self.TINY_PNG)
        source_md5 = hashlib.md5(source.read_bytes()).hexdigest()
        old_image = f"assets/{asset_dir.name}/image1.png"
        (output_root / "Mailbox.md").write_text(
            "Source: Mailbox.docx\n"
            f"Source MD5: {source_md5}\n\n"
            f"![image1]({old_image})\n",
            encoding="utf-8",
            newline="\n",
        )
        if shared:
            (output_root / "Shared.md").write_text(
                f"![shared]({old_image})\n",
                encoding="utf-8",
                newline="\n",
            )
        return root, source_root, output_root, asset_dir

    def _convert(self, build: Callable[[Document], None]) -> str:
        temporary_root = (
            Path(__file__).resolve().parents[5]
            / "build"
            / "gdd-sync"
            / "test-tmp"
        )
        temporary_root.mkdir(parents=True, exist_ok=True)
        root = temporary_root / uuid4().hex
        root.mkdir()
        try:
            source_root = root / "source"
            output_root = root / "output"
            source_root.mkdir()
            output_root.mkdir()
            source = source_root / "Mailbox.docx"
            document = Document()
            build(document)
            document.save(source)

            convert_docx_file(
                source_root,
                "Mailbox.docx",
                output_root,
            )
            return (output_root / "Mailbox.md").read_text(
                encoding="utf-8"
            )
        finally:
            shutil.rmtree(root)


if __name__ == "__main__":
    unittest.main()
