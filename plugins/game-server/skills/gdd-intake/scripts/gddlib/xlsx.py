from __future__ import annotations

from contextlib import contextmanager
from datetime import date, datetime, time
import hashlib
from itertools import zip_longest
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterator

from gddlib.common import (
    OperationResult,
    SourceFile,
    copy_file_atomic,
    md5_file,
    output_path_for,
    resolve_source_file,
    sha256_file,
    validate_output_destination,
    write_text_atomic,
)


def _validate_provenance_source(value: str) -> None:
    if "\r" in value or "\n" in value:
        raise ValueError("source path contains a line break")


@contextmanager
def _source_snapshot(
    source: SourceFile,
    staging_root: Path,
    change_error: str,
) -> Iterator[tuple[SourceFile, str]]:
    _validate_provenance_source(source.relative.as_posix())
    with NamedTemporaryFile(
        dir=staging_root.resolve(strict=True),
        prefix=".gdd-xlsx-",
        suffix=source.suffix,
        delete=False,
    ) as handle:
        snapshot_path = Path(handle.name)
    try:
        try:
            copy_file_atomic(source.absolute, snapshot_path)
        except RuntimeError as error:
            raise RuntimeError(change_error) from error
        digest = sha256_file(snapshot_path)
        yield (
            SourceFile(
                source.root,
                snapshot_path,
                source.relative,
                source.suffix,
            ),
            digest,
        )
    finally:
        snapshot_path.unlink(missing_ok=True)


def _assert_source_unchanged(
    source: SourceFile,
    snapshot: SourceFile,
    digest: str,
    change_error: str,
) -> None:
    if (
        sha256_file(snapshot.absolute) != digest
        or sha256_file(source.absolute) != digest
    ):
        raise RuntimeError(change_error)


def render_xlsx(source: SourceFile) -> tuple[str, dict[str, Any]]:
    _validate_provenance_source(source.relative.as_posix())
    try:
        import openpyxl
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
    except ImportError as error:  # pragma: no cover - environment-dependent
        raise RuntimeError(
            "openpyxl is required; install requirements.txt"
        ) from error

    workbook = openpyxl.load_workbook(
        source.absolute,
        data_only=False,
        read_only=False,
    )
    sections = [
        f"Source: {source.relative.as_posix()}",
        f"Source MD5: {md5_file(source.absolute)}",
        "",
        f"# {source.relative.stem}",
    ]
    sheet_details: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            bounds = _worksheet_bounds(worksheet)
            if bounds is None:
                continue
            min_row, min_col, max_row, max_col = bounds
            merged_values: dict[tuple[int, int], Any] = {}
            relevant_merges: list[str] = []
            for merged in worksheet.merged_cells.ranges:
                if (
                    merged.max_row < min_row
                    or merged.min_row > max_row
                    or merged.max_col < min_col
                    or merged.min_col > max_col
                ):
                    continue
                relevant_merges.append(str(merged))
                top_left = _excel_cell_value(
                    worksheet.cell(merged.min_row, merged.min_col)
                )
                for row in range(merged.min_row, merged.max_row + 1):
                    for column in range(
                        merged.min_col,
                        merged.max_col + 1,
                    ):
                        merged_values[(row, column)] = top_left

            values: list[list[str]] = []
            populated = formulas = 0
            for row_number in range(min_row, max_row + 1):
                row_values: list[str] = []
                for column_number in range(min_col, max_col + 1):
                    cell = worksheet.cell(row_number, column_number)
                    value = merged_values.get(
                        (row_number, column_number),
                        _excel_cell_value(cell),
                    )
                    if (
                        not isinstance(cell, MergedCell)
                        and cell.value is not None
                    ):
                        populated += 1
                        if cell.data_type == "f":
                            formulas += 1
                    rendered = _markdown_cell(value)
                    if (
                        cell.hyperlink
                        and cell.hyperlink.target
                        and rendered
                    ):
                        rendered = (
                            f"[{rendered}](<{cell.hyperlink.target}>)"
                        )
                    row_values.append(rendered)
                values.append(row_values)

            range_label = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{max_row}"
            )
            sections.extend(
                [
                    "",
                    f"## Sheet: {worksheet.title}",
                    "",
                    (
                        f"Range: `{range_label}` · "
                        f"Rows: {max_row - min_row + 1} · "
                        f"Columns: {max_col - min_col + 1}"
                    ),
                ]
            )
            if worksheet.sheet_state != "visible":
                sections.extend(
                    ["", f"Sheet state: `{worksheet.sheet_state}`"]
                )
            if relevant_merges:
                quoted = ", ".join(
                    f"`{item}`" for item in relevant_merges
                )
                sections.extend(
                    [
                        "",
                        (
                            f"Source merges: {quoted}. Merged labels are "
                            "repeated for Markdown readability."
                        ),
                    ]
                )

            width = max_col - min_col + 1
            sections.extend(
                [
                    "",
                    "| " + " | ".join(values[0]) + " |",
                    "| " + " | ".join(["---"] * width) + " |",
                ]
            )
            sections.extend(
                "| " + " | ".join(row) + " |" for row in values[1:]
            )
            sheet_details.append(
                {
                    "name": worksheet.title,
                    "state": worksheet.sheet_state,
                    "range": range_label,
                    "populated_cells": populated,
                    "formulas": formulas,
                    "merged_ranges": relevant_merges,
                }
            )
    finally:
        workbook.close()

    text = "\n".join(sections).rstrip() + "\n"
    details = {
        "worksheets": [sheet["name"] for sheet in sheet_details],
        "worksheet_states": {
            sheet["name"]: sheet["state"] for sheet in sheet_details
        },
        "ranges": {
            sheet["name"]: sheet["range"] for sheet in sheet_details
        },
        "populated_cells": sum(
            sheet["populated_cells"] for sheet in sheet_details
        ),
        "formulas": sum(sheet["formulas"] for sheet in sheet_details),
        "merged_ranges": {
            sheet["name"]: sheet["merged_ranges"]
            for sheet in sheet_details
        },
        "sheets": sheet_details,
    }
    return text, details


def convert_xlsx_file(
    gdd_dir: Path,
    source: str,
    output_root: Path,
) -> OperationResult:
    resolved = resolve_source_file(
        gdd_dir,
        source,
        frozenset({".xlsx"}),
    )
    output = validate_output_destination(
        output_root,
        output_path_for(resolved, output_root),
    )
    change_error = "source changed during XLSX conversion"
    with _source_snapshot(resolved, output_root, change_error) as (
        snapshot,
        source_digest,
    ):
        text, details = render_xlsx(snapshot)
        source_md5 = md5_file(snapshot.absolute)
        _assert_source_unchanged(
            resolved,
            snapshot,
            source_digest,
            change_error,
        )
    write_text_atomic(output, text)
    result_details = dict(details)
    result_details["output_sha256"] = sha256_file(output)
    return OperationResult(
        source=resolved.relative.as_posix(),
        output=output.relative_to(output_root.resolve()).as_posix(),
        status="converted",
        source_md5=source_md5,
        details=result_details,
    )


def verify_xlsx_file(
    gdd_dir: Path,
    source: str,
    output_root: Path,
) -> tuple[dict[str, Any], list[str]]:
    resolved = resolve_source_file(
        gdd_dir,
        source,
        frozenset({".xlsx"}),
    )
    output = validate_output_destination(
        output_root,
        output_path_for(resolved, output_root),
    )
    change_error = "source changed during XLSX verification"
    with _source_snapshot(resolved, output_root, change_error) as (
        snapshot,
        source_digest,
    ):
        expected_text, details = render_xlsx(snapshot)
        _assert_source_unchanged(
            resolved,
            snapshot,
            source_digest,
            change_error,
        )
    with output.open(
        "r",
        encoding="utf-8",
        errors="strict",
        newline="",
    ) as stream:
        actual_text = stream.read()
    matched = actual_text == expected_text
    errors = (
        []
        if matched
        else [
            (
                f"{resolved.relative.as_posix()}: "
                "Markdown differs from source rendering"
            )
        ]
    )
    summary = dict(details)
    summary.update(
        {
            "output": output.relative_to(
                output_root.resolve()
            ).as_posix(),
            "matched": matched,
            "expected_sha256": _sha256_text(expected_text),
            "actual_sha256": _sha256_text(actual_text),
            "first_difference": _first_difference(
                expected_text,
                actual_text,
            ),
        }
    )
    return summary, errors


def _worksheet_bounds(
    worksheet: Any,
) -> tuple[int, int, int, int] | None:
    populated_cells = [
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    coordinates = [(cell.row, cell.column) for cell in populated_cells]
    if not coordinates:
        return None

    from openpyxl.utils.cell import range_boundaries

    for cell in populated_cells:
        formula_range = getattr(cell.value, "ref", None)
        if not formula_range:
            continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(
                formula_range
            )
        except (TypeError, ValueError):
            continue
        coordinates.extend(
            [(min_row, min_col), (max_row, max_col)]
        )
    for merged in worksheet.merged_cells.ranges:
        top_left = worksheet.cell(merged.min_row, merged.min_col)
        if top_left.value is not None:
            coordinates.extend(
                [
                    (merged.min_row, merged.min_col),
                    (merged.max_row, merged.max_col),
                ]
            )
    rows = [row for row, _ in coordinates]
    columns = [column for _, column in coordinates]
    return min(rows), min(columns), max(rows), max(columns)


def _excel_cell_value(cell: Any) -> Any:
    value = cell.value
    if (
        getattr(cell, "data_type", None) != "f"
        or isinstance(value, str)
    ):
        return value
    text = getattr(value, "text", None)
    if text:
        formula_range = getattr(value, "ref", None)
        if formula_range:
            return f"{text} [array ref: {formula_range}]"
        return text
    try:
        attributes = ", ".join(
            f"{key}={item}" for key, item in value
        )
    except TypeError:
        attributes = ""
    formula_type = getattr(value, "t", type(value).__name__)
    return f"={formula_type}({attributes})"


def _markdown_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        rendered = value.isoformat(sep=" ")
    elif isinstance(value, (date, time)):
        rendered = value.isoformat()
    elif isinstance(value, bool):
        rendered = "TRUE" if value else "FALSE"
    else:
        rendered = str(value)
    return (
        rendered.replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\r", "<br>")
        .replace("\n", "<br>")
    )


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _first_difference(
    expected_text: str,
    actual_text: str,
) -> dict[str, Any] | None:
    expected_lines = expected_text.splitlines(keepends=True)
    actual_lines = actual_text.splitlines(keepends=True)
    for line, (expected, actual) in enumerate(
        zip_longest(expected_lines, actual_lines),
        start=1,
    ):
        if expected != actual:
            return {
                "line": line,
                "expected": expected,
                "actual": actual,
            }
    return None
