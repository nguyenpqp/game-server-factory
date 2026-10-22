from __future__ import annotations

import hashlib
import os
import re
import shutil
import stat
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence
from uuid import uuid4

from gddlib import markdown as markdown_parser
from gddlib.common import (
    OperationResult,
    SourceFile,
    copy_file_atomic,
    md5_file,
    output_path_for,
    resolve_source_file,
    sha256_file,
    write_text_atomic,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
WP_NS = (
    "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
)
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"

NS = {
    "w": W_NS,
    "r": R_NS,
    "a": A_NS,
    "v": V_NS,
    "wp": WP_NS,
    "pic": PIC_NS,
}

SUPPORTED_MC_REQUIREMENT_NAMESPACES = {
    W_NS,
    R_NS,
    A_NS,
    V_NS,
    WP_NS,
    PIC_NS,
}


def _validate_provenance_source(value: str) -> None:
    if "\r" in value or "\n" in value:
        raise ValueError("source path contains a line break")


def _slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", ascii_value).strip("-").lower()
    return slug or "document"


def _markdown_cell(value: Any) -> str:
    if value is None:
        return ""
    rendered = str(value)
    return (
        rendered.replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\r", "<br>")
        .replace("\n", "<br>")
    )


def _markdown_alt(value: str, in_table: bool = False) -> str:
    normalized = re.sub(r"\s+", " ", value).strip()
    if in_table:
        return (
            normalized.replace("\\", "&#92;")
            .replace("[", "&#91;")
            .replace("]", "&#93;")
        )
    return (
        normalized.replace("\\", "\\\\")
        .replace("[", "\\[")
        .replace("]", "\\]")
    )


MARKDOWN_PUNCTUATION_RE = re.compile(
    r"[!#&()*+\-<>\[\]\\_`{|}~]"
)
WORD_TOC_BOOKMARK_RE = re.compile(r"_Toc\d+\Z")
HYPERLINK_FIELD_RE = re.compile(
    r"\bHYPERLINK\b\s+(?:\"([^\"]+)\"|(\S+))",
    re.IGNORECASE,
)
HYPERLINK_FIELD_ANCHOR_RE = re.compile(
    r"\\l\s+(?:\"([^\"]+)\"|(\S+))",
    re.IGNORECASE,
)


def _markdown_text(value: str, in_table: bool) -> str:
    if in_table:
        return MARKDOWN_PUNCTUATION_RE.sub(
            lambda match: f"&#{ord(match.group(0))};",
            value,
        )
    escaped = MARKDOWN_PUNCTUATION_RE.sub(
        lambda match: f"\\{match.group(0)}",
        value,
    )
    return re.sub(r"(?<=\d)\.(?=\s)", r"\\.", escaped)


def _markdown_heading_slug(value: str) -> str:
    return markdown_parser.github_heading_slug(value)


def _xml_local_name(element: Any) -> str:
    tag = element.tag
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _xpath_nodes(element: Any, expression: str) -> list[Any]:
    try:
        return list(element.xpath(expression, namespaces=NS))
    except TypeError:
        return list(element.xpath(expression))


def _iter_word_blocks(element: Any) -> Iterator[Any]:
    for child in element.iterchildren():
        local = _xml_local_name(child)
        if local in {"p", "tbl"}:
            yield child
        elif local in {
            "sdt",
            "sdtContent",
            "customXml",
            "ins",
            "moveTo",
        }:
            yield from _iter_word_blocks(child)


def _iter_visible_run_nodes(element: Any) -> Iterator[Any]:
    for child in element.iterchildren():
        local = _xml_local_name(child)
        if local == "AlternateContent":
            alternatives = list(child.iterchildren())
            selected = next(
                (
                    alternative
                    for alternative in alternatives
                    if (
                        _xml_local_name(alternative) == "Choice"
                        and all(
                            alternative.nsmap.get(prefix)
                            in SUPPORTED_MC_REQUIREMENT_NAMESPACES
                            for prefix in alternative.get(
                                "Requires",
                                "",
                            ).split()
                        )
                    )
                ),
                None,
            )
            if selected is None:
                selected = next(
                    (
                        alternative
                        for alternative in alternatives
                        if _xml_local_name(alternative) == "Fallback"
                    ),
                    None,
                )
            if selected is not None:
                yield from _iter_visible_run_nodes(selected)
            continue
        yield child
        if local not in {"drawing", "pict"}:
            yield from _iter_visible_run_nodes(child)


def _media_relative_path(part_name: str) -> Path:
    normalized = part_name.replace("\\", "/")
    parts = tuple(normalized.split("/"))
    windows_devices = {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        "CLOCK$",
        "CONIN$",
        "CONOUT$",
        *(f"COM{suffix}" for suffix in "123456789¹²³"),
        *(f"LPT{suffix}" for suffix in "123456789¹²³"),
    }
    if len(parts) >= 3 and parts[:2] == ("word", "media"):
        media_parts = parts[2:]
    elif len(parts) >= 2 and parts[0] == "media":
        media_parts = parts[1:]
    else:
        media_parts = ()
    if (
        normalized.startswith("/")
        or not media_parts
        or any(
            part in {"", ".", ".."}
            or part != part.rstrip(" .")
            or any(ord(character) < 32 for character in part)
            or any(character in '<>:"|?*' for character in part)
            or (
                part.split(".", 1)[0].rstrip(" ").upper()
                in windows_devices
            )
            for part in media_parts
        )
    ):
        raise RuntimeError(f"unsafe DOCX media part: {part_name}")
    return Path(*parts)


def _media_extraction_target(asset_dir: Path, relative: Path) -> Path:
    asset_root = asset_dir.resolve(strict=True)
    target = (asset_dir / relative).resolve()
    if target == asset_root or not target.is_relative_to(asset_root):
        raise RuntimeError(
            f"unsafe DOCX media target: {relative.as_posix()}"
        )
    return target


def _installed_media_name(
    index: int,
    original: Path,
    payload: bytes,
) -> Path:
    digest = hashlib.sha256(payload).hexdigest()
    return Path(
        f"image_{index}_{digest[:12]}{original.suffix.lower()}"
    )


class DocxRenderer:
    def __init__(
        self,
        source: Path,
        output: Path,
        source_label: str,
        source_digest: str | None = None,
    ) -> None:
        _validate_provenance_source(source_label)
        try:
            from docx import Document
        except ImportError as error:  # pragma: no cover - environment-dependent
            raise RuntimeError(
                "python-docx is required; install requirements.txt"
            ) from error

        self.source = source
        self.output = output
        self.source_label = source_label
        self.document = Document(source)
        self.style_elements = {
            style.get(f"{{{W_NS}}}styleId"): style
            for style in self.document.styles.element.iterchildren()
            if (
                _xml_local_name(style) == "style"
                and style.get(f"{{{W_NS}}}styleId")
            )
        }
        digest = source_digest or sha256_file(source)
        self.asset_slug = f"{_slugify(source.stem)}-{digest[:12]}"
        self.asset_dir = output.parent / "assets" / self.asset_slug
        self.asset_rel = Path("assets") / self.asset_slug
        self.extracted: dict[str, Path] = {}
        self.image_occurrences = 0
        self.hyperlinks = 0
        self.paragraphs = 0
        self.tables = 0
        self.blocks = 0
        self.toc_heading_targets = self._toc_heading_targets()
        self.word_toc_targets = self._word_toc_heading_targets()
        self._extract_media()

    def _extract_media(self) -> None:
        self.asset_dir.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(self.source) as archive:
            media: list[tuple[Path, bytes]] = []
            for member in archive.namelist():
                normalized_member = member.replace("\\", "/")
                if (
                    not normalized_member.startswith(
                        ("word/media/", "media/")
                    )
                    or normalized_member.endswith("/")
                ):
                    continue
                relative = _media_relative_path(member)
                payload = archive.read(member)
                media.append((relative, payload))
            media.sort(key=lambda item: item[0].as_posix())
            for index, (relative, payload) in enumerate(media, start=1):
                installed = _installed_media_name(
                    index,
                    relative,
                    payload,
                )
                target = _media_extraction_target(
                    self.asset_dir,
                    installed,
                )
                target.parent.mkdir(parents=True, exist_ok=True)
                if not target.exists() or target.read_bytes() != payload:
                    target.write_bytes(payload)
                if hashlib.sha256(target.read_bytes()).digest() != (
                    hashlib.sha256(payload).digest()
                ):
                    raise RuntimeError(
                        f"media hash mismatch after extraction: {member}"
                    )
                self.extracted[relative.as_posix()] = target

    def _relationship_target(self, relationship_id: str) -> Any | None:
        return self.document.part.rels.get(relationship_id)

    def _image_token(
        self,
        relationship_id: str,
        container: Any,
        in_table: bool = False,
    ) -> str:
        relationship = self._relationship_target(relationship_id)
        if relationship is None:
            raise RuntimeError(
                f"missing image relationship: {relationship_id}"
            )
        if getattr(relationship, "is_external", False):
            raise RuntimeError(
                "linked external image is unsupported: "
                f"{relationship_id}"
            )
        try:
            relative = _media_relative_path(
                str(relationship.target_part.partname).lstrip("/")
            )
        except (AttributeError, KeyError, ValueError) as error:
            raise RuntimeError(
                f"unresolved image relationship: {relationship_id}"
            ) from error
        target = self.extracted.get(relative.as_posix())
        if target is None:
            try:
                blob = relationship.target_part.blob
            except (AttributeError, KeyError) as error:
                raise RuntimeError(
                    f"unresolved image relationship: {relationship_id}"
                ) from error
            installed = _installed_media_name(
                len(self.extracted) + 1,
                relative,
                blob,
            )
            target = _media_extraction_target(
                self.asset_dir,
                installed,
            )
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(blob)
            self.extracted[relative.as_posix()] = target

        alt = ""
        for node in _xpath_nodes(
            container,
            ".//wp:docPr | .//pic:cNvPr",
        ):
            alt = (
                node.get("descr")
                or node.get("title")
                or node.get("name")
                or ""
            )
            if alt:
                break
        if not alt or re.fullmatch(
            r"Picture\s+\d+",
            alt,
            re.IGNORECASE,
        ):
            alt = target.stem
        self.image_occurrences += 1
        markdown_target = (
            self.asset_rel / target.relative_to(self.asset_dir)
        ).as_posix()
        return (
            f"![{_markdown_alt(alt, in_table)}]({markdown_target})"
        )

    @staticmethod
    def _property_value(properties: Any, name: str) -> bool | None:
        if properties is None:
            return None
        element = properties.find(f"{{{W_NS}}}{name}")
        if element is None:
            return None
        value = element.get(f"{{{W_NS}}}val")
        return value is None or value.lower() not in {
            "0",
            "false",
            "off",
            "no",
        }

    def _style_chain(self, style_id: str | None) -> Iterator[Any]:
        visited: set[str] = set()
        while style_id and style_id not in visited:
            visited.add(style_id)
            style = self.style_elements.get(style_id)
            if style is None:
                return
            yield style
            based_on = style.find(f"{{{W_NS}}}basedOn")
            style_id = (
                based_on.get(f"{{{W_NS}}}val")
                if based_on is not None
                else None
            )

    def _run_property_values(self, run_xml: Any) -> dict[str, bool]:
        run_properties = run_xml.find(f"{{{W_NS}}}rPr")
        names = ("b", "i", "strike", "vanish", "webHidden")
        values: dict[str, bool] = {}

        def inherit(properties: Any) -> None:
            for name in names:
                if name in values:
                    continue
                value = self._property_value(properties, name)
                if value is not None:
                    values[name] = value

        inherit(run_properties)
        run_style = (
            run_properties.find(f"{{{W_NS}}}rStyle")
            if run_properties is not None
            else None
        )
        run_style_id = (
            run_style.get(f"{{{W_NS}}}val")
            if run_style is not None
            else None
        )
        for style in self._style_chain(run_style_id):
            inherit(style.find(f"{{{W_NS}}}rPr"))

        defaults = self.document.styles.element.find(
            f"{{{W_NS}}}docDefaults/{{{W_NS}}}rPrDefault/"
            f"{{{W_NS}}}rPr"
        )
        inherit(defaults)
        return values

    def _run_style_markers(self, run_xml: Any) -> tuple[str, ...]:
        values = self._run_property_values(run_xml)
        return tuple(
            marker
            for name, marker in (
                ("b", "**"),
                ("i", "*"),
                ("strike", "~~"),
            )
            if values.get(name, False)
        )

    def _run_style_delimiters(self, run_xml: Any) -> tuple[str, str]:
        markers = self._run_style_markers(run_xml)
        return "".join(markers), "".join(reversed(markers))

    def _run_is_hidden(self, run_xml: Any) -> bool:
        values = self._run_property_values(run_xml)
        return values.get("vanish", False) or values.get(
            "webHidden",
            False,
        )

    def _run_style(self, run_xml: Any, text: str) -> str:
        if not text:
            return text
        prefix, suffix = self._run_style_delimiters(run_xml)
        if not prefix:
            return text
        content_start = len(text) - len(text.lstrip())
        content_end = len(text.rstrip())
        if content_start == content_end:
            return text
        return (
            f"{text[:content_start]}"
            f"{prefix}{text[content_start:content_end]}{suffix}"
            f"{text[content_end:]}"
        )

    def _render_run(
        self,
        run_xml: Any,
        in_table: bool = False,
    ) -> str:
        if self._run_is_hidden(run_xml):
            return ""

        pieces: list[str] = []
        text_pieces: list[str] = []

        def flush_text() -> None:
            if text_pieces:
                pieces.append(
                    self._run_style(run_xml, "".join(text_pieces))
                )
                text_pieces.clear()

        for node in _iter_visible_run_nodes(run_xml):
            local = _xml_local_name(node)
            if node.tag == f"{{{W_NS}}}t":
                text_pieces.append(
                    _markdown_text(node.text or "", in_table)
                )
            elif node.tag == f"{{{W_NS}}}tab":
                text_pieces.append("\t")
            elif node.tag in {
                f"{{{W_NS}}}br",
                f"{{{W_NS}}}cr",
            }:
                text_pieces.append("<br>")
            elif node.tag == f"{{{W_NS}}}noBreakHyphen":
                text_pieces.append("‑")
            elif node.tag == f"{{{W_NS}}}softHyphen":
                text_pieces.append("\u00ad")
            elif local == "drawing":
                flush_text()
                for blip in _xpath_nodes(
                    node,
                    ".//a:blip[@r:embed or @r:link]",
                ):
                    embedded = blip.get(f"{{{R_NS}}}embed")
                    linked = blip.get(f"{{{R_NS}}}link")
                    if linked is not None:
                        raise RuntimeError(
                            "linked external image is unsupported: "
                            f"{linked}"
                        )
                    if embedded:
                        pieces.append(
                            self._image_token(
                                embedded,
                                node,
                                in_table,
                            )
                        )
            elif local == "pict":
                flush_text()
                for image_data in _xpath_nodes(
                    node,
                    ".//v:imagedata[@r:id]",
                ):
                    relationship_id = image_data.get(f"{{{R_NS}}}id")
                    if relationship_id:
                        pieces.append(
                            self._image_token(
                                relationship_id,
                                node,
                                in_table,
                            )
                        )
        flush_text()
        return "".join(pieces)

    def _visible_heading_text(self, paragraph_xml: Any) -> str:
        parts: list[str] = []

        def collect(parent: Any) -> None:
            for child in parent.iterchildren():
                local = _xml_local_name(child)
                if local == "r":
                    if self._run_is_hidden(child):
                        continue
                    for node in _iter_visible_run_nodes(child):
                        if node.tag == f"{{{W_NS}}}t":
                            parts.append(node.text or "")
                        elif node.tag == f"{{{W_NS}}}tab":
                            parts.append("\t")
                        elif node.tag == f"{{{W_NS}}}noBreakHyphen":
                            parts.append("‑")
                        elif node.tag == f"{{{W_NS}}}softHyphen":
                            parts.append("\u00ad")
                elif local in {
                    "hyperlink",
                    "smartTag",
                    "sdt",
                    "sdtContent",
                    "ins",
                    "moveTo",
                    "fldSimple",
                    "customXml",
                }:
                    collect(child)

        collect(paragraph_xml)
        return "".join(parts)

    @staticmethod
    def _unique_heading_slug(
        base_slug: str,
        occurrences: dict[str, int],
    ) -> str:
        slug = base_slug
        while slug in occurrences:
            occurrence = occurrences.get(base_slug, 0) + 1
            occurrences[base_slug] = occurrence
            slug = f"{base_slug}-{occurrence}"
        occurrences.setdefault(base_slug, 0)
        occurrences.setdefault(slug, 0)
        return slug

    def _word_toc_heading_targets(self) -> dict[str, str]:
        from docx.text.paragraph import Paragraph

        targets: dict[str, str] = {}
        slug_occurrences: dict[str, int] = {}
        for block in _iter_word_blocks(self.document.element.body):
            if _xml_local_name(block) != "p":
                continue
            paragraph = Paragraph(block, self.document)
            if self._heading_level(paragraph) is None:
                continue
            visible_text = self._visible_heading_text(block)
            base_slug = _markdown_heading_slug(visible_text)
            slug = self._unique_heading_slug(
                base_slug,
                slug_occurrences,
            )
            for bookmark in _xpath_nodes(
                block,
                ".//w:bookmarkStart",
            ):
                name = bookmark.get(f"{{{W_NS}}}name")
                if name and WORD_TOC_BOOKMARK_RE.fullmatch(name):
                    targets[name] = slug
        return targets

    def _toc_heading_targets(self) -> tuple[tuple[str, str], ...]:
        from docx.text.paragraph import Paragraph

        targets: list[tuple[str, str]] = []
        slug_occurrences: dict[str, int] = {}
        for block in _iter_word_blocks(self.document.element.body):
            if _xml_local_name(block) != "p":
                continue
            paragraph = Paragraph(block, self.document)
            if self._heading_level(paragraph) is None:
                continue
            visible_text = self._visible_heading_text(block).strip()
            if not visible_text:
                continue
            slug = self._unique_heading_slug(
                _markdown_heading_slug(visible_text),
                slug_occurrences,
            )
            targets.append((visible_text, slug))
        return tuple(targets)

    @staticmethod
    def _toc_style_level(paragraph: Any) -> int | None:
        if paragraph.style is None:
            return None
        values = {
            paragraph.style.name or "",
            paragraph.style.style_id or "",
        }
        for value in values:
            normalized = re.sub(r"[_-]+", " ", value).strip()
            match = re.fullmatch(
                r"(?:toc|table of contents)\s*([1-6])",
                normalized,
                re.IGNORECASE,
            )
            if match is not None:
                return int(match.group(1))
        return None

    @staticmethod
    def _toc_title(value: str) -> str:
        label = value.split("\t", 1)[0].strip()
        match = re.match(
            r"^\d+(?:\.\d+)*(?:\.)?\s+(.+?)\s*$",
            label,
        )
        return (match.group(1) if match is not None else label).strip()

    @staticmethod
    def _toc_title_key(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", value).casefold()
        return re.sub(r"\s+", " ", normalized).strip()

    @staticmethod
    def _toc_titles_match(toc_title: str, heading_title: str) -> bool:
        if toc_title == heading_title:
            return True
        separators = (" — ", " – ", " - ", ": ")
        return any(
            toc_title.startswith(heading_title + separator)
            or heading_title.startswith(toc_title + separator)
            for separator in separators
        )

    def _toc_heading_target(
        self,
        value: str,
    ) -> tuple[str, str] | None:
        toc_title = self._toc_title_key(self._toc_title(value))
        matches = [
            (heading, slug)
            for heading, slug in self.toc_heading_targets
            if self._toc_titles_match(
                toc_title,
                self._toc_title_key(self._toc_title(heading)),
            )
        ]
        return matches[0] if len(matches) == 1 else None

    def _anchor_target(self, anchor: str) -> str | None:
        if WORD_TOC_BOOKMARK_RE.fullmatch(anchor):
            heading_target = self.word_toc_targets.get(anchor)
            return f"#{heading_target}" if heading_target else None
        return f"#{anchor}"

    def _field_hyperlink_target(
        self,
        instruction: str,
    ) -> str | None:
        if re.search(r"\bHYPERLINK\b", instruction, re.IGNORECASE) is None:
            return None
        target_match = HYPERLINK_FIELD_RE.search(instruction)
        target = (
            next(
                (
                    value
                    for value in target_match.groups()
                    if value is not None
                ),
                None,
            )
            if target_match is not None
            else None
        )
        if target is not None and target.startswith("\\"):
            target = None
        anchor_match = HYPERLINK_FIELD_ANCHOR_RE.search(instruction)
        anchor = (
            next(
                (
                    value
                    for value in anchor_match.groups()
                    if value is not None
                ),
                None,
            )
            if anchor_match is not None
            else None
        )
        if target and anchor:
            return f"{target}#{anchor}"
        if target:
            return target
        if anchor:
            return self._anchor_target(anchor)
        return None

    def _render_link(self, label: str, target: str | None) -> str:
        label_start = len(label) - len(label.lstrip())
        label_end = len(label.rstrip())
        label_text = label[label_start:label_end]
        if not target or not label_text:
            return label
        self.hyperlinks += 1
        return (
            f"{label[:label_start]}"
            f"[{label_text}](<{target}>)"
            f"{label[label_end:]}"
        )

    @staticmethod
    def _field_character_type(run_xml: Any) -> str | None:
        field = run_xml.find(f"{{{W_NS}}}fldChar")
        return (
            field.get(f"{{{W_NS}}}fldCharType")
            if field is not None
            else None
        )

    @classmethod
    def _complex_field(
        cls,
        children: Sequence[Any],
        start: int,
    ) -> tuple[str, list[Any], int] | None:
        depth = 0
        separated = False
        instruction_parts: list[str] = []
        result: list[Any] = []
        for index in range(start, len(children)):
            child = children[index]
            field_type = (
                cls._field_character_type(child)
                if _xml_local_name(child) == "r"
                else None
            )
            if field_type == "begin":
                depth += 1
                if depth > 1 and separated:
                    result.append(child)
                elif depth == 1:
                    instruction_parts.extend(
                        node.text or ""
                        for node in _xpath_nodes(
                            child,
                            ".//w:instrText",
                        )
                    )
                continue
            if depth == 0:
                return None
            if field_type == "separate" and depth == 1:
                separated = True
                continue
            if field_type == "end":
                depth -= 1
                if depth == 0:
                    return "".join(instruction_parts), result, index
                if separated:
                    result.append(child)
                continue
            if separated:
                result.append(child)
            elif depth == 1:
                instruction_parts.extend(
                    node.text or ""
                    for node in _xpath_nodes(child, ".//w:instrText")
                )
        return None

    def _render_inline_elements(
        self,
        elements: Iterable[Any],
        in_table: bool,
        *,
        strip: bool,
    ) -> str:
        parts: list[str] = []
        previous_run_markers: tuple[str, ...] | None = None

        def render_children(children: Iterable[Any]) -> None:
            nonlocal previous_run_markers
            nodes = list(children)
            index = 0
            while index < len(nodes):
                child = nodes[index]
                local = _xml_local_name(child)
                if (
                    local == "r"
                    and self._field_character_type(child) == "begin"
                ):
                    field = self._complex_field(nodes, index)
                    if field is not None:
                        instruction, result, end = field
                        previous_run_markers = None
                        target = self._field_hyperlink_target(
                            instruction
                        )
                        if target:
                            label = self._render_inline_elements(
                                result,
                                in_table,
                                strip=False,
                            )
                            parts.append(self._render_link(label, target))
                        else:
                            render_children(result)
                        index = end + 1
                        continue
                if local == "r":
                    rendered = self._render_run(child, in_table)
                    if rendered:
                        if rendered.isspace():
                            if parts:
                                parts[-1] += rendered
                            else:
                                parts.append(rendered)
                            index += 1
                            continue
                        markers = self._run_style_markers(child)
                        prefix = "".join(markers)
                        suffix = "".join(reversed(markers))
                        shared = 0
                        if previous_run_markers is not None:
                            for previous, current in zip(
                                previous_run_markers,
                                markers,
                            ):
                                if previous != current:
                                    break
                                shared += 1
                        previous_suffix = (
                            "".join(reversed(previous_run_markers))
                            if previous_run_markers
                            else ""
                        )
                        previous_part = parts[-1] if parts else ""
                        previous_core = previous_part.rstrip()
                        previous_whitespace = previous_part[
                            len(previous_core) :
                        ]
                        current_core = rendered.lstrip()
                        current_whitespace = rendered[
                            : len(rendered) - len(current_core)
                        ]
                        if (
                            shared
                            and "![" not in rendered
                            and parts
                            and previous_core.endswith(previous_suffix)
                            and current_core.startswith(prefix)
                        ):
                            close_previous = "".join(
                                reversed(
                                    previous_run_markers[shared:]
                                )
                            )
                            open_current = "".join(markers[shared:])
                            parts[-1] = (
                                previous_core[: -len(previous_suffix)]
                                + close_previous
                                + previous_whitespace
                                + current_whitespace
                                + open_current
                                + current_core[len(prefix) :]
                            )
                        else:
                            parts.append(rendered)
                        previous_run_markers = (
                            None if "![" in rendered else markers
                        )
                elif local == "hyperlink":
                    previous_run_markers = None
                    label = self._render_inline_elements(
                        child.iterchildren(),
                        in_table,
                        strip=False,
                    )
                    relationship_id = child.get(f"{{{R_NS}}}id")
                    anchor = child.get(f"{{{W_NS}}}anchor")
                    target: str | None = None
                    if relationship_id:
                        relationship = self._relationship_target(
                            relationship_id
                        )
                        if relationship is not None:
                            target = str(relationship.target_ref)
                    elif anchor:
                        target = self._anchor_target(anchor)
                    parts.append(self._render_link(label, target))
                elif local == "fldSimple":
                    instruction = child.get(f"{{{W_NS}}}instr", "")
                    target = self._field_hyperlink_target(instruction)
                    if target:
                        previous_run_markers = None
                        label = self._render_inline_elements(
                            child.iterchildren(),
                            in_table,
                            strip=False,
                        )
                        parts.append(self._render_link(label, target))
                    else:
                        render_children(child.iterchildren())
                elif local == "bookmarkStart":
                    name = child.get(f"{{{W_NS}}}name")
                    if (
                        name
                        and name != "_GoBack"
                        and WORD_TOC_BOOKMARK_RE.fullmatch(name) is None
                    ):
                        previous_run_markers = None
                        parts.append(f'<a id="{name}"></a>')
                elif local in {
                    "smartTag",
                    "sdt",
                    "sdtContent",
                    "ins",
                    "moveTo",
                    "customXml",
                }:
                    render_children(child.iterchildren())
                index += 1

        render_children(elements)
        rendered = "".join(parts)
        if strip:
            rendered = rendered.strip()
        if not in_table:
            rendered = re.sub(
                r"(?<=\d)\.(?=\s)",
                r"\\.",
                rendered,
            )
        return rendered

    def _render_inline_container(
        self,
        container: Any,
        in_table: bool = False,
        *,
        strip: bool = True,
    ) -> str:
        return self._render_inline_elements(
            container.iterchildren(),
            in_table,
            strip=strip,
        )

    def _list_properties(
        self,
        paragraph: Any,
    ) -> tuple[str, int] | None:
        numbering_properties = None
        paragraph_properties = paragraph._p.find(f"{{{W_NS}}}pPr")
        if paragraph_properties is not None:
            numbering_properties = paragraph_properties.find(
                f"{{{W_NS}}}numPr"
            )
        if (
            numbering_properties is None
            and paragraph.style is not None
        ):
            style_properties = paragraph.style._element.find(
                f"{{{W_NS}}}pPr"
            )
            if style_properties is not None:
                numbering_properties = style_properties.find(
                    f"{{{W_NS}}}numPr"
                )

        style_name = (
            paragraph.style.name if paragraph.style is not None else ""
        ) or ""
        style_lower = style_name.lower().replace(" ", "")
        level = 0
        numbering_id: int | None = None
        if numbering_properties is not None:
            level_element = numbering_properties.find(
                f"{{{W_NS}}}ilvl"
            )
            numbering_element = numbering_properties.find(
                f"{{{W_NS}}}numId"
            )
            if level_element is not None:
                level = int(
                    level_element.get(f"{{{W_NS}}}val", "0")
                )
            if numbering_element is not None:
                numbering_id = int(
                    numbering_element.get(f"{{{W_NS}}}val", "0")
                )

        if numbering_id == 0:
            return None
        if "listbullet" in style_lower:
            return "-", level
        if "listnumber" in style_lower:
            return "1.", level
        if numbering_id is None:
            return None

        kind = "-"
        try:
            numbering = self.document.part.numbering_part.element
            numbering_nodes = _xpath_nodes(
                numbering,
                f"./w:num[@w:numId='{numbering_id}']",
            )
            if numbering_nodes:
                abstract_id_nodes = _xpath_nodes(
                    numbering_nodes[0],
                    "./w:abstractNumId",
                )
                abstract_id = (
                    abstract_id_nodes[0].get(f"{{{W_NS}}}val")
                    if abstract_id_nodes
                    else None
                )
                abstract_nodes = _xpath_nodes(
                    numbering,
                    (
                        "./w:abstractNum"
                        f"[@w:abstractNumId='{abstract_id}']"
                    ),
                )
                if abstract_nodes:
                    levels = _xpath_nodes(
                        abstract_nodes[0],
                        f"./w:lvl[@w:ilvl='{level}']",
                    )
                    formats = (
                        _xpath_nodes(levels[0], "./w:numFmt")
                        if levels
                        else []
                    )
                    format_name = (
                        formats[0].get(f"{{{W_NS}}}val")
                        if formats
                        else "bullet"
                    )
                    if format_name == "none":
                        return None
                    if format_name != "bullet":
                        kind = "1."
        except (
            AttributeError,
            IndexError,
            KeyError,
            TypeError,
            ValueError,
        ):
            pass
        return kind, level

    def _heading_level(self, paragraph: Any) -> int | None:
        paragraph_properties = paragraph._p.find(f"{{{W_NS}}}pPr")

        def outline_level(properties: Any) -> int | None:
            if properties is None:
                return None
            outline = properties.find(f"{{{W_NS}}}outlineLvl")
            if outline is None:
                return None
            try:
                value = int(outline.get(f"{{{W_NS}}}val", "9"))
            except ValueError:
                return None
            return value + 1 if 0 <= value <= 5 else 0

        direct_level = outline_level(paragraph_properties)
        if direct_level is not None:
            return direct_level or None

        paragraph_style = (
            paragraph_properties.find(f"{{{W_NS}}}pStyle")
            if paragraph_properties is not None
            else None
        )
        style_id = (
            paragraph_style.get(f"{{{W_NS}}}val")
            if paragraph_style is not None
            else getattr(paragraph.style, "style_id", None)
        )
        for style in self._style_chain(style_id):
            inherited_level = outline_level(
                style.find(f"{{{W_NS}}}pPr")
            )
            if inherited_level is not None:
                return inherited_level or None

        name = (
            paragraph.style.name if paragraph.style is not None else ""
        ) or ""
        match = re.search(
            r"heading\s*([1-6])",
            name,
            re.IGNORECASE,
        )
        if match:
            return int(match.group(1))
        if name.lower() in {"title", "document title"}:
            return 1
        return None

    def render_paragraph(
        self,
        paragraph: Any,
        in_cell: bool = False,
    ) -> str:
        self.paragraphs += 1
        if self._toc_style_level(paragraph) is not None:
            target = self._toc_heading_target(
                self._visible_heading_text(paragraph._p)
            )
            if target is None:
                return ""
            heading, slug = target
            return self._render_link(
                _markdown_text(heading, False),
                f"#{slug}",
            )
        content = self._render_inline_container(
            paragraph._p,
            in_cell,
        )
        if not content:
            return ""
        list_properties = self._list_properties(paragraph)
        if in_cell:
            if list_properties:
                marker, level = list_properties
                indent = "&nbsp;" * (level * 4)
                return f"{indent}{marker} {content}".replace(
                    "\n",
                    "<br>",
                )
            return content.replace("\n", "<br>")
        heading = self._heading_level(paragraph)
        if heading:
            return f"{'#' * heading} {content}"
        if list_properties:
            marker, level = list_properties
            return f"{'    ' * level}{marker} {content}"
        return content

    def _cell_text(self, cell: Any) -> str:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        parts: list[str] = []
        for block in _iter_word_blocks(cell._tc):
            local = _xml_local_name(block)
            if local == "p":
                rendered = self.render_paragraph(
                    Paragraph(block, cell),
                    in_cell=True,
                )
                if rendered:
                    parts.append(rendered)
            elif local == "tbl":
                nested = Table(block, cell)
                seen_cells: set[Any] = set()
                for row in nested.rows:
                    row_parts: list[str] = []
                    for item in row.cells:
                        if item._tc in seen_cells:
                            continue
                        seen_cells.add(item._tc)
                        row_parts.append(self._cell_text(item))
                    if row_parts:
                        parts.append(" / ".join(row_parts))
        return _markdown_cell("<br>".join(parts)).replace(
            "&lt;br&gt;",
            "<br>",
        )

    def render_table(self, table: Any) -> str:
        self.tables += 1
        seen_cells: set[Any] = set()
        rows: list[list[str]] = []
        for row in table.rows:
            rendered_row: list[str] = []
            for cell in row.cells:
                if cell._tc in seen_cells:
                    rendered_row.append("")
                    continue
                seen_cells.add(cell._tc)
                rendered_row.append(self._cell_text(cell))
            rows.append(rendered_row)
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        normalized = [
            row + [""] * (width - len(row))
            for row in rows
        ]

        def row_line(values: Sequence[str]) -> str:
            return "| " + " | ".join(values) + " |"

        lines = [
            row_line(normalized[0]),
            row_line(["---"] * width),
        ]
        lines.extend(row_line(row) for row in normalized[1:])
        return "\n".join(lines)

    def render(self) -> tuple[str, dict[str, Any]]:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        rendered_blocks: list[str] = []
        for block in _iter_word_blocks(self.document.element.body):
            local = _xml_local_name(block)
            if local == "p":
                rendered = self.render_paragraph(
                    Paragraph(block, self.document)
                )
            elif local == "tbl":
                rendered = self.render_table(
                    Table(block, self.document)
                )
            else:  # pragma: no cover - guarded by _iter_word_blocks
                continue
            if rendered:
                rendered_blocks.append(rendered)
                self.blocks += 1

        source = self.source_label
        source_md5 = md5_file(self.source)
        body = "\n\n".join(rendered_blocks)
        text = (
            (
                f"Source: {source}\n"
                f"Source MD5: {source_md5}\n\n"
                f"{body}"
            )
            if body
            else (
                f"Source: {source}\n"
                f"Source MD5: {source_md5}\n\n"
            )
        )
        details = {
            "blocks": self.blocks,
            "paragraphs": self.paragraphs,
            "tables": self.tables,
            "hyperlinks": self.hyperlinks,
            "image_occurrences": self.image_occurrences,
            "unique_media": len(self.extracted),
            "asset_directory": self.asset_dir.as_posix(),
        }
        return text, details


def convert_docx(
    source: Path,
    output: Path,
    source_label: str,
    source_digest: str,
) -> dict[str, Any]:
    renderer = DocxRenderer(
        source,
        output,
        source_label,
        source_digest,
    )
    text, details = renderer.render()
    if text.endswith("\n\n"):
        output.write_text(text, encoding="utf-8", newline="\n")
    else:
        write_text_atomic(output, text)
    return details


def _is_reparse_point(path: Path) -> bool:
    try:
        metadata = os.lstat(path)
    except FileNotFoundError:
        return False
    attributes = getattr(metadata, "st_file_attributes", 0)
    reparse_flag = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0)
    return stat.S_ISLNK(metadata.st_mode) or bool(
        attributes & reparse_flag
    )


def _absolute_path(path: Path) -> Path:
    return Path(os.path.abspath(path))


def _assert_no_output_reparse(path: Path) -> None:
    absolute = _absolute_path(path)
    for candidate in (*reversed(absolute.parents), absolute):
        if _is_reparse_point(candidate):
            raise RuntimeError(
                "output path contains a symlink or reparse point: "
                f"{candidate}"
            )


def _validate_output_path(
    output_root: Path,
    output: Path,
) -> tuple[Path, Path]:
    lexical_root = _absolute_path(output_root)
    lexical_output = _absolute_path(output)
    _assert_no_output_reparse(lexical_root)
    _assert_no_output_reparse(lexical_output)
    root = lexical_root.resolve()
    resolved_parent = lexical_output.parent.resolve()
    if not resolved_parent.is_relative_to(root):
        raise RuntimeError(
            "resolved output parent is outside output root: "
            f"{resolved_parent}"
        )
    if lexical_output.exists() and lexical_output.resolve() != (
        resolved_parent / lexical_output.name
    ):
        raise RuntimeError(
            f"resolved output is outside output root: {lexical_output}"
        )
    return root, lexical_output


def _staged_output_root(source: SourceFile, output: Path) -> Path:
    root = _absolute_path(output)
    relative_output = source.relative.with_suffix(".md")
    for _ in relative_output.parts:
        root = root.parent
    return root


def _assert_plain_directory_tree(root: Path) -> None:
    if _is_reparse_point(root):
        raise RuntimeError(
            f"asset directory is a symlink or reparse point: {root}"
        )
    if not root.is_dir():
        raise RuntimeError(f"asset path is not a directory: {root}")
    pending = [root]
    while pending:
        current = pending.pop()
        with os.scandir(current) as entries:
            for entry in entries:
                path = Path(entry.path)
                if _is_reparse_point(path):
                    raise RuntimeError(
                        "asset directory contains a symlink or reparse "
                        f"point: {path}"
                    )
                metadata = entry.stat(follow_symlinks=False)
                if stat.S_ISDIR(metadata.st_mode):
                    pending.append(path)
                elif not stat.S_ISREG(metadata.st_mode):
                    raise RuntimeError(
                        f"asset directory contains a special file: {path}"
                    )


def directory_fingerprint(root: Path) -> list[tuple[str, str]]:
    _assert_plain_directory_tree(root)
    files: list[Path] = []
    pending = [root]
    while pending:
        current = pending.pop()
        with os.scandir(current) as entries:
            for entry in entries:
                path = Path(entry.path)
                metadata = entry.stat(follow_symlinks=False)
                if stat.S_ISDIR(metadata.st_mode):
                    pending.append(path)
                elif stat.S_ISREG(metadata.st_mode):
                    files.append(path)
    return sorted(
        (
            path.relative_to(root).as_posix(),
            sha256_file(path),
        )
        for path in files
    )


def _install_missing_versioned_assets(
    staged: Path,
    installed: Path,
) -> tuple[Path, ...]:
    _assert_plain_directory_tree(staged)
    _assert_plain_directory_tree(installed)
    installed_root = installed.resolve(strict=True)
    added: list[Path] = []
    try:
        for relative, digest in directory_fingerprint(staged):
            source = staged / Path(relative)
            target = installed / Path(relative)
            if not target.resolve().is_relative_to(installed_root):
                raise RuntimeError(
                    f"versioned asset target escapes directory: {target}"
                )
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists() or target.is_symlink():
                if (
                    _is_reparse_point(target)
                    or not target.is_file()
                    or sha256_file(target) != digest
                ):
                    raise RuntimeError(
                        "versioned asset file differs from source: "
                        f"{target}"
                    )
                continue
            try:
                copy_file_atomic(source, target)
            except BaseException:
                if target.is_file() and not _is_reparse_point(target):
                    target.unlink()
                raise
            added.append(target)
        return tuple(added)
    except BaseException:
        _remove_installed_asset_files(added)
        raise


def _remove_installed_asset_files(files: Iterable[Path]) -> None:
    for path in reversed(tuple(files)):
        if _is_reparse_point(path) or not path.is_file():
            raise RuntimeError(
                f"installed asset changed before rollback: {path}"
            )
        path.unlink()


def _remove_obsolete_versioned_assets(
    installed: Path,
    staged: Path,
) -> None:
    expected = directory_fingerprint(staged)
    expected_paths = {relative for relative, _ in expected}
    _assert_plain_directory_tree(installed)
    for relative, _ in directory_fingerprint(installed):
        if relative in expected_paths:
            continue
        target = installed / Path(relative)
        if _is_reparse_point(target) or not target.is_file():
            raise RuntimeError(
                f"obsolete asset changed before deletion: {target}"
            )
        target.unlink()
    directories = sorted(
        (path for path in installed.rglob("*") if path.is_dir()),
        key=lambda path: len(path.parts),
        reverse=True,
    )
    for directory in directories:
        if not any(directory.iterdir()):
            directory.rmdir()
    if directory_fingerprint(installed) != expected:
        raise RuntimeError(
            "versioned asset directory differs after migration: "
            f"{installed}"
        )


def _was_referenced_by_owner(
    directory: Path,
    old_directories: Iterable[Path],
) -> bool:
    resolved = directory.resolve(strict=True)
    return any(
        candidate.resolve(strict=True) == resolved
        for candidate in old_directories
    )


def _remove_installed_assets_after_failed_output(
    asset_directory: Path,
) -> None:
    _assert_plain_directory_tree(asset_directory)
    shutil.rmtree(asset_directory)


def convert_docx_staged(
    source: SourceFile,
    output: Path,
    output_root: Path | None = None,
    old_asset_directories: Iterable[Path] = (),
) -> dict[str, Any]:
    _validate_provenance_source(source.relative.as_posix())
    root, output = _validate_output_path(
        output_root or _staged_output_root(source, output),
        output,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    staging_root = output.parent / (
        f".{_slugify(output.stem)}-{uuid4().hex}.gdd-tool"
    )
    staging_root.mkdir(mode=0o777 if os.name == "nt" else 0o700)
    old_asset_directories = tuple(old_asset_directories)
    installed_assets: Path | None = None
    installed_asset_files: tuple[Path, ...] = ()
    migrate_same_version = False
    try:
        source_snapshot = staging_root / source.absolute.name
        try:
            copy_file_atomic(source.absolute, source_snapshot)
        except RuntimeError as error:
            raise RuntimeError(
                "source changed during DOCX conversion"
            ) from error
        source_digest = sha256_file(source_snapshot)
        staged_output = staging_root / output.name
        details = convert_docx(
            source_snapshot,
            staged_output,
            source.relative.as_posix(),
            source_digest,
        )
        details["_source_md5"] = md5_file(source_snapshot)
        if (
            sha256_file(source_snapshot) != source_digest
            or sha256_file(source.absolute) != source_digest
        ):
            raise RuntimeError(
                "source changed during DOCX conversion"
            )

        staged_assets = Path(details["asset_directory"])
        final_assets_parent = output.parent / "assets"
        if _is_reparse_point(final_assets_parent):
            raise RuntimeError(
                "asset destination is a symlink or reparse point: "
                f"{final_assets_parent}"
            )
        final_assets = final_assets_parent / staged_assets.name
        stale_assets = _validate_stale_docx_assets(
            old_asset_directories,
            final_assets,
            output,
            root,
            require_new=False,
        )
        final_assets.parent.mkdir(parents=True, exist_ok=True)
        if final_assets.exists() or final_assets.is_symlink():
            if _is_reparse_point(final_assets):
                raise RuntimeError(
                    "asset destination is a symlink or reparse point: "
                    f"{final_assets}"
                )
            if not final_assets.is_dir():
                raise RuntimeError(
                    "asset destination is not a directory: "
                    f"{final_assets}"
                )
            if directory_fingerprint(
                final_assets
            ) != directory_fingerprint(staged_assets):
                if not _was_referenced_by_owner(
                    final_assets,
                    old_asset_directories,
                ):
                    raise RuntimeError(
                        "versioned asset directory differs from source: "
                        f"{final_assets}"
                    )
                installed_asset_files = _install_missing_versioned_assets(
                    staged_assets,
                    final_assets,
                )
                migrate_same_version = True
        else:
            staged_assets.replace(final_assets)
            installed_assets = final_assets

        try:
            if output.exists():
                if not output.is_file():
                    raise RuntimeError(
                        f"output exists but is not a file: {output}"
                    )
                output.chmod(stat.S_IREAD | stat.S_IWRITE)
            staged_output.replace(output)
        except BaseException:
            if installed_assets is not None:
                _remove_installed_assets_after_failed_output(
                    installed_assets
                )
            if installed_asset_files:
                _remove_installed_asset_files(installed_asset_files)
            raise

        if migrate_same_version:
            _remove_obsolete_versioned_assets(
                final_assets,
                staged_assets,
            )

        details["asset_directory"] = final_assets.resolve(
            strict=True
        ).as_posix()
        details["output_sha256"] = sha256_file(output)
        removed = _delete_validated_stale_assets(stale_assets)
        details["removed_asset_directories"] = [
            path.as_posix() for path in removed
        ]
        return details
    finally:
        shutil.rmtree(staging_root, ignore_errors=True)


def _asset_root_for_target(target: Path) -> Path | None:
    directory = target.parent
    for candidate in (directory, *directory.parents):
        if candidate.parent.name.casefold() == "assets":
            return candidate
    return None


def asset_directories_referenced_by(markdown: Path) -> set[Path]:
    text = markdown_parser._read_markdown(markdown)
    markdown_parser._reject_invalid_text(text, markdown)
    _, visible_lines, _ = markdown_parser._markdown_scan_lines(text)
    expected_assets = (markdown.parent / "assets").resolve()
    directories: set[Path] = set()
    for kind in ("image", "link"):
        targets_by_line = markdown_parser._targets_by_line(
            visible_lines,
            images=kind == "image",
        )
        for targets in targets_by_line.values():
            for target in targets:
                if target.error is not None:
                    continue
                try:
                    relative_target = markdown_parser._local_target(
                        target.target,
                        kind=kind,
                    )
                except (UnicodeError, ValueError):
                    continue
                if relative_target is None:
                    continue
                resolved = (
                    markdown.parent / relative_target
                ).resolve()
                asset = _asset_root_for_target(resolved)
                if (
                    asset is None
                    or asset.parent != expected_assets
                    or re.fullmatch(
                        (
                            r"[a-z0-9]+(?:-[a-z0-9]+)*-"
                            r"[0-9a-f]{12}"
                        ),
                        asset.name,
                    )
                    is None
                    or not resolved.is_file()
                ):
                    continue
                directories.add(asset)
    return directories


def _active_markdown_files(output_root: Path) -> Iterator[Path]:
    pending = [output_root]
    while pending:
        current = pending.pop()
        children = sorted(
            current.iterdir(),
            key=lambda path: path.name.casefold(),
            reverse=True,
        )
        for child in children:
            relative = child.relative_to(output_root)
            if relative.parts[0] in {"_deleted_", "instruction"}:
                continue
            if _is_reparse_point(child):
                raise RuntimeError(
                    "active Markdown scope contains a symlink or "
                    f"reparse point: {child}"
                )
            if child.is_dir():
                pending.append(child)
            elif child.is_file() and child.suffix.lower() == ".md":
                yield child


def active_asset_references(
    output_root: Path,
    excluding: Path | None = None,
) -> set[Path]:
    root = output_root.resolve()
    excluded = excluding.resolve() if excluding is not None else None
    references: set[Path] = set()
    for markdown in _active_markdown_files(root):
        if excluded is not None and markdown.resolve() == excluded:
            continue
        references.update(asset_directories_referenced_by(markdown))
    return references


def _validate_stale_docx_assets(
    old_dirs: Iterable[Path],
    new_dir: Path,
    owner_output: Path,
    output_root: Path,
    *,
    require_new: bool,
) -> tuple[Path, ...]:
    root = output_root.resolve(strict=True)
    owner = _absolute_path(owner_output)
    if _is_reparse_point(owner):
        raise RuntimeError(
            f"owner output is a symlink or reparse point: {owner}"
        )
    if owner.exists() and not owner.is_file():
        raise RuntimeError(f"owner output is not a file: {owner}")
    resolved_owner = owner.resolve()
    if not resolved_owner.is_relative_to(root):
        raise RuntimeError(
            f"owner output is outside output root: {owner_output}"
        )

    expected_parent = resolved_owner.parent / "assets"
    if _is_reparse_point(expected_parent):
        raise RuntimeError(
            "owner assets is a symlink or reparse point: "
            f"{expected_parent}"
        )
    expected_parent = expected_parent.resolve()

    if not new_dir.is_absolute():
        raise RuntimeError(
            f"new asset directory is not absolute: {new_dir}"
        )
    if _is_reparse_point(new_dir):
        raise RuntimeError(
            f"new asset directory is a symlink or reparse point: {new_dir}"
        )
    resolved_new = new_dir.resolve()
    if resolved_new.parent != expected_parent:
        raise RuntimeError(
            f"new asset directory is outside owner assets: {resolved_new}"
        )
    if require_new:
        _assert_plain_directory_tree(resolved_new)

    candidates: list[Path] = []
    current_version: list[Path] = []
    for raw_candidate in sorted(
        set(old_dirs),
        key=lambda path: path.as_posix(),
    ):
        if not raw_candidate.is_absolute():
            raise RuntimeError(
                "old asset directory is not absolute: "
                f"{raw_candidate}"
            )
        if _is_reparse_point(raw_candidate):
            raise RuntimeError(
                "old asset directory is a symlink or reparse point: "
                f"{raw_candidate}"
            )
        try:
            candidate = raw_candidate.resolve(strict=True)
        except FileNotFoundError as error:
            raise RuntimeError(
                f"old asset directory does not exist: {raw_candidate}"
            ) from error
        if raw_candidate != candidate:
            raise RuntimeError(
                "old asset directory is unresolved or traverses a "
                f"symlink: {raw_candidate}"
            )
        if candidate.parent != expected_parent:
            raise RuntimeError(
                "old asset directory is outside owner assets: "
                f"{candidate}"
            )
        if candidate == resolved_new:
            _assert_plain_directory_tree(candidate)
            current_version.append(candidate)
            continue
        _assert_plain_directory_tree(candidate)
        candidates.append(candidate)

    referenced_candidates = candidates + current_version
    if referenced_candidates:
        active_references = active_asset_references(
            root,
            excluding=resolved_owner,
        )
        for candidate in referenced_candidates:
            if candidate in active_references:
                raise RuntimeError(
                    "old asset directory is still referenced: "
                    f"{candidate}"
                )
    return tuple(candidates)


def _delete_validated_stale_assets(
    candidates: Iterable[Path],
) -> tuple[Path, ...]:
    removed: list[Path] = []
    for candidate in candidates:
        if (
            _is_reparse_point(candidate)
            or candidate.resolve(strict=True) != candidate
        ):
            raise RuntimeError(
                "old asset directory changed before deletion: "
                f"{candidate}"
            )
        _assert_plain_directory_tree(candidate)
        shutil.rmtree(candidate)
        removed.append(candidate)
    return tuple(removed)


def remove_stale_docx_assets(
    old_dirs: Iterable[Path],
    new_dir: Path,
    owner_output: Path,
    output_root: Path,
) -> tuple[Path, ...]:
    candidates = _validate_stale_docx_assets(
        old_dirs,
        new_dir,
        owner_output,
        output_root,
        require_new=True,
    )
    return _delete_validated_stale_assets(candidates)


def convert_docx_file(
    gdd_dir: Path,
    source: str,
    output_root: Path,
) -> OperationResult:
    resolved = resolve_source_file(
        gdd_dir,
        source,
        frozenset({".docx"}),
    )
    lexical_root = _absolute_path(output_root)
    _assert_no_output_reparse(lexical_root)
    root = lexical_root.resolve()
    output = output_path_for(resolved, root)
    root, output = _validate_output_path(root, output)
    previous_assets = (
        asset_directories_referenced_by(output)
        if output.exists()
        else set()
    )
    details = convert_docx_staged(
        resolved,
        output,
        root,
        previous_assets,
    )
    source_md5 = details.pop("_source_md5")
    return OperationResult(
        source=resolved.relative.as_posix(),
        output=output.relative_to(root).as_posix(),
        status="converted",
        source_md5=source_md5,
        details=details,
    )
